# -*- coding: utf-8 -*-
"""Regression tests for security boundaries and honest QA reporting."""
from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import pytest

from conftest import REPO_ROOT, SCRIPTS


FIGURE_SCRIPTS = REPO_ROOT / "skills" / "academic-figure" / "scripts"


def load_script(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize(
    "source",
    [
        r"\small 摘要",
        r"\fontsize{10pt}{11pt}\selectfont 摘要",
        r"\resizebox{\textwidth}{!}{摘要}",
        r"\vspace{-1em} 摘要",
        r"\linespread{0.8}\selectfont 摘要",
        r"\newpage 摘要",
        r"\baselineskip=10pt 摘要",
        r"\parskip = 0pt 摘要",
        r"\hspace{-1em} 摘要",
        r"\kern -1em 摘要",
        r"\hskip-1em 摘要",
        r"\vskip -1em 摘要",
        r"\makebox[0pt]{重叠摘要}",
        r"\makebox[0\linewidth]{重叠摘要}",
        r"\smash{重叠摘要}",
        r"\raisebox{-1em}{重叠摘要}",
    ],
)
def test_abstract_layout_rejects_squeezing_or_manual_pagination(source):
    audit_tex = load_script("hardening_abstract_layout", SCRIPTS / "audit_tex.py")
    assert audit_tex.abstract_layout_violations(source)


def test_abstract_layout_allows_semantic_emphasis_and_unbreakable_keywords():
    audit_tex = load_script("hardening_abstract_layout_valid", SCRIPTS / "audit_tex.py")
    source = (
        r"\textbf{关键模型与 14.248 dBZ}\keywords{\mbox{强对流临近预报}}"
        r"\hspace{1em}\vspace{1ex}\kern 1pt\hskip 1pt\vskip 1pt"
        r"\makebox[2cm]{正常宽度}\raisebox{1pt}{正常抬升}"
    )
    assert audit_tex.abstract_layout_violations(source) == []


@pytest.mark.parametrize(
    "source",
    [
        r"\iffalse\small\fi 正常摘要",
        r"\begin{verbatim}\small\end{verbatim} 正常摘要",
        r"\verb|\small| 正常摘要",
    ],
)
def test_abstract_layout_ignores_inactive_or_literal_commands(source):
    audit_tex = load_script("hardening_abstract_layout_inactive", SCRIPTS / "audit_tex.py")
    assert audit_tex.abstract_layout_violations(source) == []


def test_abstract_layout_fails_closed_on_unresolved_conditionals():
    audit_tex = load_script("hardening_abstract_layout_unresolved", SCRIPTS / "audit_tex.py")
    violations = audit_tex.abstract_layout_violations(r"\ifdraft\small\fi 摘要")
    assert "UNRESOLVED_TEX_CONDITIONAL" in violations


@pytest.mark.parametrize(
    "source",
    [
        r"\newgeometry{margin=1cm}",
        r"\setlength{\textheight}{30cm}",
        r"\addtolength{\textwidth}{1cm}",
        r"\oddsidemargin = 0pt",
        r"\textheight 30cm",
        r"\evensidemargin=0pt",
        r"\topmargin = -1cm",
        r"\advance\headsep by 1pt",
    ],
)
def test_page_layout_rejects_geometry_and_page_dimension_overrides(source):
    audit_tex = load_script("hardening_page_layout_override", SCRIPTS / "audit_tex.py")
    assert audit_tex.page_layout_overrides(source)


def test_page_layout_allows_local_spacing_and_inactive_or_literal_examples():
    audit_tex = load_script("hardening_page_layout_valid", SCRIPTS / "audit_tex.py")
    source = (
        r"\setlength{\tabcolsep}{4pt}\addtolength{\parskip}{1pt}"
        "\n" + r"\setlength{\figurewidth}{\dimexpr\textwidth-2cm\relax}"
        r"\iffalse\newgeometry{margin=1cm}\fi"
        r"\verb|\setlength{\textheight}{30cm}|"
    )
    assert audit_tex.page_layout_overrides(source) == []


def test_r_png_device_quotes_untrusted_output_path():
    compose = load_script("hardening_compose", FIGURE_SCRIPTS / "compose.py")
    malicious = 'panel"); system("whoami"); #\\next.png'
    expected_literal = json.dumps(malicious.replace("\\", "/"), ensure_ascii=False)
    call = compose.r_png_device({"width_mm": 65, "height_mm": 40, "dpi": 300}, malicious)
    assert call.startswith(f"png({expected_literal}, ")
    assert "eval(parse" not in call


def test_markdown_migration_rejects_missing_abstract_and_source_escape(tmp_path):
    migration = load_script("hardening_migration", SCRIPTS / "migrate_markdown_to_tex.py")
    source = tmp_path / "legacy"
    source.mkdir()
    (source / "chapter.md").write_text("# 正文\n内容", encoding="utf-8")
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps({"chapters": [{"title": "正文", "content_file": "chapter.md"}]}, ensure_ascii=False),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="abstract"):
        migration.migrate(manifest, source, tmp_path / "tex")

    outside = tmp_path / "outside.md"
    outside.write_text("outside", encoding="utf-8")
    manifest.write_text(
        json.dumps(
            {"abstract": "摘要", "chapters": [{"title": "正文", "content_file": "../outside.md"}]},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="source-dir"):
        migration.migrate(manifest, source, tmp_path / "tex")


def test_markdown_migration_rejects_case_insensitive_output_collisions(tmp_path):
    migration = load_script("hardening_migration_collision", SCRIPTS / "migrate_markdown_to_tex.py")
    source = tmp_path / "legacy"
    source.mkdir()
    (source / "A.md").write_text("A", encoding="utf-8")
    (source / "a.md").write_text("a", encoding="utf-8")
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "abstract": "摘要",
                "chapters": [
                    {"title": "A", "content_file": "A.md"},
                    {"title": "a", "content_file": "a.md"},
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="冲突"):
        migration.migrate(manifest, source, tmp_path / "tex")


def test_word_figure_path_stays_inside_project(tmp_path):
    build_docx = load_script("hardening_build_docx", SCRIPTS / "build_docx.py")
    image = tmp_path / "figure.png"
    image.write_bytes(b"not decoded in this path-only test")
    assert build_docx.resolve_figure_path(tmp_path, "figure.png") == image.resolve()

    outside = tmp_path.parent / "outside.png"
    outside.write_bytes(b"outside")
    with pytest.raises(ValueError, match="工作目录"):
        build_docx.resolve_figure_path(tmp_path, "../outside.png")

    vector = tmp_path / "figure.svg"
    vector.write_text("<svg/>", encoding="utf-8")
    with pytest.raises(ValueError, match="不支持"):
        build_docx.resolve_figure_path(tmp_path, "figure.svg")


def test_word_derivative_renders_production_figure_wrappers_in_order(tmp_path):
    import fitz
    from docx import Document

    build_docx = load_script("hardening_build_docx_wrappers", SCRIPTS / "build_docx.py")
    figures = tmp_path / "figures"
    figures.mkdir()
    pdf = fitz.open()
    page = pdf.new_page(width=300, height=180)
    page.insert_text((36, 90), "RESULT FIGURE")
    pdf.save(figures / "result.pdf")
    pdf.close()

    document = Document()
    build_docx.configure_styles(document)
    source = (
        "上一图结论已经解释。\n"
        "\\evidencefigure\n"
        "  {result}\n"
        "  {中文结果图题}\n"
        "  {fig:result}"
        "\n下一段继续分析误差趋势。"
    )
    build_docx.add_tex_content(
        document, source, tmp_path, skip_first_section=False
    )
    assert len(document.inline_shapes) == 1
    visible = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "上一图结论已经解释" in visible
    assert "中文结果图题" in visible
    assert "下一段继续分析误差趋势" in visible
    assert "fig:result" not in visible
    with pytest.raises(FileNotFoundError, match="图片不存在"):
        build_docx.add_tex_content(
            Document(),
            r"\frameworkfigure{missing}{中文缺失图}{fig:missing}",
            tmp_path,
            skip_first_section=False,
        )


def test_word_manifest_identity_check_does_not_reject_subject_words(tmp_path):
    build_docx = load_script("hardening_build_docx_identity", SCRIPTS / "build_docx.py")
    audit_docx = load_script("hardening_audit_docx_identity", SCRIPTS / "audit_docx.py")
    for legitimate in ("学校食堂排队优化", "高校实验室设备调度", "team formation optimization"):
        assert build_docx.IDENTITY_RE.search(legitimate) is None
        assert audit_docx.IDENTITY_RE.search(legitimate) is None
    for identity_field in (
        "学校名称：清华大学",
        "队员：张三",
        "队员 2：李四",
        "指导老师：王五",
        "School Name: Example University",
        "Team ID: 12345",
    ):
        assert build_docx.IDENTITY_RE.search(identity_field)
        assert audit_docx.IDENTITY_RE.search(identity_field)
    (tmp_path / "abstract.tex").write_text("摘要", encoding="utf-8")
    (tmp_path / "chapter.tex").write_text("正文", encoding="utf-8")
    manifest = tmp_path / "paper.json"
    payload = {
        "title": "学校食堂排队优化",
        "keywords": ["team formation optimization", "排队优化", "设备调度"],
        "abstract_tex_path": "abstract.tex",
        "appendix_pseudocode": {
            "required": False,
            "reason": "本测试仅检查纯解析文本，不涉及程序求解或数值实验。",
        },
        "chapters": [
            {
                "chapter_id": "problem",
                "title": "高校实验室设备调度",
                "role": "problem",
                "order": 1,
                "tex_path": "chapter.tex",
            }
        ],
    }
    manifest.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    assert build_docx.read_input(manifest)["title"] == "学校食堂排队优化"

    payload["title"] = "学校：示例大学"
    manifest.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ValueError, match="身份信息"):
        build_docx.read_input(manifest)


def test_word_tex_reader_expands_inline_static_inputs_but_not_comments(tmp_path):
    build_docx = load_script("hardening_build_docx_inline_input", SCRIPTS / "build_docx.py")
    (tmp_path / "sub.tex").write_text("中间内容", encoding="utf-8")
    (tmp_path / "extra.tex").write_text("补充内容", encoding="utf-8")
    source = tmp_path / "main.tex"
    source.write_text(
        "前缀 \\input{sub} 后缀\n"
        "% \\input{missing}\n"
        "继续 \\include{extra}。\n",
        encoding="utf-8",
    )
    expanded = build_docx.read_tex_source(source, tmp_path)
    assert "前缀 中间内容 后缀" in expanded
    assert "% \\input{missing}" in expanded
    assert "继续 补充内容。" in expanded


def test_repository_does_not_bundle_unverified_word_template():
    template = REPO_ROOT / "assets" / "paper-template" / "研赛论文Word标准模板.docx"
    assert not template.exists(), "旧 v2.1 社区 Word 稿不得冒充 2026 官方模板随仓库发布"


def test_external_local_file_relationships_are_detected_and_scrubbed(tmp_path):
    audit_docx = load_script("hardening_external_rel_audit", SCRIPTS / "audit_docx.py")
    build_docx = load_script("hardening_external_rel_scrub", SCRIPTS / "build_docx.py")
    unsafe_targets = [
        (r"\\server\share\secret.dotm", "External"),
        ("//server/share/secret.dotm", "External"),
        ("/home/user/secret.dotm", "External"),
        ("file%3A///C%3A/Users/Test/secret.dotm", "External"),
        ("%5C%5Cserver%5Cshare%5Csecret.dotm", "External"),
    ]
    for target, mode in unsafe_targets:
        assert audit_docx.is_external_local_file_target(target, mode)
        assert build_docx.is_external_local_file_target(target, mode)
    for target, mode in [
        ("https://example.com/template.dotm", "External"),
        ("mailto:test@example.com", "External"),
        ("/word/media/image1.png", ""),
    ]:
        assert not audit_docx.is_external_local_file_target(target, mode)
        assert not build_docx.is_external_local_file_target(target, mode)

    from docx import Document

    template = tmp_path / "blank-template.docx"
    Document().save(template)
    unsafe = tmp_path / "unsafe_external_relationship.docx"
    rel_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(template) as source, zipfile.ZipFile(unsafe, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                root = ET.fromstring(data)
                ET.SubElement(
                    root,
                    f"{{{rel_namespace}}}Relationship",
                    {
                        "Id": "rUnsafeLocalTemplate",
                        "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate",
                        "Target": r"\\server\share\secret.dotm",
                        "TargetMode": "External",
                    },
                )
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, data)

    unsafe_report = audit_docx.audit(unsafe)
    unsafe_check = next(
        check for check in unsafe_report["checks"]
        if check["code"] == "no_external_file_relationships"
    )
    assert unsafe_check["ok"] is False
    build_docx.scrub_package(unsafe)
    cleaned_report = audit_docx.audit(unsafe)
    cleaned_check = next(
        check for check in cleaned_report["checks"]
        if check["code"] == "no_external_file_relationships"
    )
    assert cleaned_check["ok"] is True


def test_word_derivative_builds_and_passes_package_audit(tmp_path):
    build_docx = load_script("hardening_build_docx_e2e", SCRIPTS / "build_docx.py")
    audit_docx = load_script("hardening_audit_docx_e2e", SCRIPTS / "audit_docx.py")
    assert build_docx.tex_table_rows("A & & C \\\\") == [["A", "", "C"]]
    assert build_docx.tex_table_rows(r"A \& B & 含义 \\") == [["A & B", "含义"]]
    compact_tabular = (
        r"\centering\begin{tabular*}{\textwidth}{lll}\toprule "
        r"符号 & 含义 & 单位 \\ \midrule "
        r"$x$ & 决策变量 & -- \\ \bottomrule\end{tabular*}"
    )
    assert build_docx.tex_table_rows(compact_tabular) == [
        ["符号", "含义", "单位"],
        ["x", "决策变量", "--"],
    ]
    longtable = r"""
\begin{longtable}{lll}
\caption{主要符号说明}\label{tab:sym} \\
\toprule
符号 & 含义 & 单位 \\
\midrule
\endfirsthead
\toprule
符号 & 含义 & 单位 \\
\midrule
\endhead
\midrule
\multicolumn{3}{r}{续下页} \\
\bottomrule
\endfoot
\bottomrule
\endlastfoot
$x$ & 决策变量 & -- \\
$c$ & 单位成本 & 元 \\
\end{longtable}
"""
    assert build_docx.tex_table_rows(longtable) == [
        ["符号", "含义", "单位"],
        ["x", "决策变量", "--"],
        ["c", "单位成本", "元"],
    ]
    assert "μ" not in build_docx.tex_to_word_text(r"\multicolumn{3}{r}{续下页}")
    (tmp_path / "abstract.tex").write_text(
        r"本研究构建 Model 1，并得到 95\% 的验证结果。", encoding="utf-8"
    )
    (tmp_path / "chapter.tex").write_text(
        r"\section{问题重述}" + "\n" + "正文包含 English 123 与可复核结论。", encoding="utf-8"
    )
    (tmp_path / "symbols.tex").write_text(
        "\n".join(
            [
                r"\section{主要符号说明}",
                r"全文主要符号见表\ref{tab:main-symbols}。",
                r"\begin{table*}[htbp]",
                r"\caption{全文主要符号、定义与单位}",
                r"\label{tab:main-symbols}",
                r"\begin{tabular*}{\textwidth}{lll}",
                r"\toprule",
                r"符号 & 定义 & 单位 \\",
                r"\midrule",
                r"$x$ & 决策变量 & -- \\",
                r"$c$ & 单位成本 & 元 \\",
                r"\bottomrule",
                r"\end{tabular*}",
                r"\end{table*}",
            ]
        ),
        encoding="utf-8",
    )
    (tmp_path / "references.tex").write_text("参考来源条目。", encoding="utf-8")
    (tmp_path / "appendix.tex").write_text("复现环境与关键参数。", encoding="utf-8")
    manifest = tmp_path / "paper.json"
    manifest.write_text(
        json.dumps(
            {
                "title": "Test 模型研究",
                "keywords": ["建模", "Model", "验证"],
                "abstract_tex_path": "abstract.tex",
                "appendix_pseudocode": {"required": True},
                "chapters": [
                    {
                        "chapter_id": "symbols",
                        "title": "主要符号说明",
                        "role": "preliminary",
                        "order": 1,
                        "level": 1,
                        "tex_path": "symbols.tex",
                    },
                    {
                        "chapter_id": "problem",
                        "title": "问题重述",
                        "role": "problem",
                        "order": 2,
                        "level": 1,
                        "tex_path": "chapter.tex",
                    },
                    {
                        "chapter_id": "refs", "title": "参考文献", "role": "references",
                        "order": 3, "level": 1, "tex_path": "references.tex",
                    },
                    {
                        "chapter_id": "appendix", "title": "附录 A 可复现代码框架", "role": "appendix",
                        "order": 4, "level": 1, "tex_path": "appendix.tex",
                    },
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    config = tmp_path / "比赛配置.json"
    config.write_text(
        json.dumps(
            {"contest": {"edition_cn": "二十三", "verified_against_official_rules": True}},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    output = tmp_path / "paper.docx"
    from docx import Document

    local_template = tmp_path / "blank-template.docx"
    Document().save(local_template)
    build_docx.build(
        SimpleNamespace(
            input=manifest,
            template=local_template,
            output=output,
            manifest_out=None,
            contest_config=config,
        )
    )
    assert Document(output).styles["Huawei Appendix Code"].font.name == "Times New Roman"
    report = audit_docx.audit(output)
    assert report["status"] == "PASS", json.dumps(report["failures"], ensure_ascii=False, indent=2)
    symbol_check = next(
        item for item in report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert symbol_check["data_row_count"] == 2
    assert symbol_check["caption_ok"] is True
    assert symbol_check["caption_tags"] == ["tab:main-symbols"]
    assert symbol_check["caption_tag_unique"] is True
    assert symbol_check["references"][0]["tags"] == ["tab:main-symbols"]

    placeholder_body = tmp_path / "paper_body_placeholder.docx"
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(placeholder_body, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                root = audit_docx.parse_xml(data)
                text_node = next(
                    node for node in root.xpath(".//w:body//w:t", namespaces=audit_docx.NS)
                    if "正文包含" in (node.text or "")
                )
                text_node.text = (text_node.text or "") + " TODO"
                data = audit_docx.etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            target.writestr(item, data)
    placeholder_report = audit_docx.audit(placeholder_body)
    placeholder_check = next(
        item for item in placeholder_report["checks"]
        if item["code"] == "template_placeholders"
    )
    assert placeholder_check["ok"] is False
    assert "TODO" in placeholder_check["matches"]

    grid_table = tmp_path / "paper_grid_symbol_table.docx"
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(grid_table, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                root = audit_docx.parse_xml(data)
                table = root.xpath(".//w:tbl", namespaces=audit_docx.NS)[0]
                inside_v = table.find("w:tblPr/w:tblBorders/w:insideV", audit_docx.NS)
                assert inside_v is not None
                inside_v.set(audit_docx.W + "val", "single")
                data = audit_docx.etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            target.writestr(item, data)
    grid_report = audit_docx.audit(grid_table)
    grid_check = next(
        item for item in grid_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert grid_check["ok"] is False
    assert any("三线表" in failure for failure in grid_check["failures"])

    missing_caption = tmp_path / "paper_missing_symbol_caption.docx"
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(missing_caption, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                root = audit_docx.parse_xml(data)
                body = root.find("w:body", audit_docx.NS)
                caption = next(
                    paragraph for paragraph in body.findall("w:p", audit_docx.NS)
                    if "符号" in audit_docx.element_text(paragraph)
                    and paragraph.xpath(
                        './w:pPr/w:pStyle[contains(translate(@w:val, "ABCDEFGHIJKLMNOPQRSTUVWXYZ", "abcdefghijklmnopqrstuvwxyz"), "huaweitablecaption")]',
                        namespaces=audit_docx.NS,
                    )
                )
                body.remove(caption)
                data = audit_docx.etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            target.writestr(item, data)
    caption_report = audit_docx.audit(missing_caption)
    caption_check = next(
        item for item in caption_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert caption_check["ok"] is False
    assert caption_check["caption_ok"] is False

    inherited_grid = tmp_path / "paper_inherited_grid_style.docx"
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(inherited_grid, "w") as target:
        styles_root = audit_docx.parse_xml(source.read("word/styles.xml"))
        grid_style = next(
            style for style in styles_root.xpath(".//w:style", namespaces=audit_docx.NS)
            if (style.xpath("./w:name/@w:val", namespaces=audit_docx.NS) or [""])[0] == "Table Grid"
        )
        grid_style_id = grid_style.get(audit_docx.W + "styleId")
        custom_style_id = "InheritedGridForAudit"
        custom_style = audit_docx.etree.Element(audit_docx.W + "style")
        custom_style.set(audit_docx.W + "type", "table")
        custom_style.set(audit_docx.W + "styleId", custom_style_id)
        custom_name = audit_docx.etree.SubElement(custom_style, audit_docx.W + "name")
        custom_name.set(audit_docx.W + "val", "Inherited Grid For Audit")
        based_on = audit_docx.etree.SubElement(custom_style, audit_docx.W + "basedOn")
        based_on.set(audit_docx.W + "val", grid_style_id)
        styles_root.append(custom_style)
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/styles.xml":
                data = audit_docx.etree.tostring(
                    styles_root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            elif item.filename == "word/document.xml":
                root = audit_docx.parse_xml(data)
                table = root.xpath(".//w:tbl", namespaces=audit_docx.NS)[0]
                properties = table.find("w:tblPr", audit_docx.NS)
                style = properties.find("w:tblStyle", audit_docx.NS)
                if style is None:
                    style = audit_docx.etree.Element(audit_docx.W + "tblStyle")
                    properties.insert(0, style)
                style.set(audit_docx.W + "val", custom_style_id)
                borders = properties.find("w:tblBorders", audit_docx.NS)
                for edge in ("left", "right", "insideH", "insideV"):
                    border = borders.find(f"w:{edge}", audit_docx.NS)
                    if border is not None:
                        borders.remove(border)
                data = audit_docx.etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            target.writestr(item, data)
    inherited_report = audit_docx.audit(inherited_grid)
    inherited_check = next(
        item for item in inherited_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert inherited_check["ok"] is False
    assert inherited_check["table_style_name"] == "Inherited Grid For Audit"
    assert "Table Grid" in inherited_check["table_style_chain"]

    def write_docx_mutation(target_path, mutate):
        with zipfile.ZipFile(output) as source, zipfile.ZipFile(target_path, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "word/document.xml":
                    root = audit_docx.parse_xml(data)
                    mutate(root)
                    data = audit_docx.etree.tostring(
                        root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                target.writestr(item, data)

    mismatched_tag = tmp_path / "paper_mismatched_symbol_tag.docx"

    def mismatch_reference_tag(root):
        reference = next(
            paragraph for paragraph in root.xpath(".//w:body/w:p", namespaces=audit_docx.NS)
            if "全文主要符号见表" in audit_docx.element_text(paragraph)
        )
        for text_node in reference.xpath(".//w:t", namespaces=audit_docx.NS):
            text_node.text = (text_node.text or "").replace(
                "[tab:main-symbols]", "[tab:wrong-symbols]"
            )

    write_docx_mutation(mismatched_tag, mismatch_reference_tag)
    mismatch_report = audit_docx.audit(mismatched_tag)
    mismatch_check = next(
        item for item in mismatch_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert mismatch_check["ok"] is False
    assert mismatch_check["reference_before_table"] is False

    missing_caption_tag = tmp_path / "paper_missing_symbol_caption_tag.docx"

    def remove_caption_tag(root):
        caption = next(
            paragraph for paragraph in root.xpath(".//w:body/w:p", namespaces=audit_docx.NS)
            if "[tab:main-symbols]" in audit_docx.element_text(paragraph)
            and paragraph.xpath(
                './w:pPr/w:pStyle[contains(translate(@w:val, "ABCDEFGHIJKLMNOPQRSTUVWXYZ", "abcdefghijklmnopqrstuvwxyz"), "huaweitablecaption")]',
                namespaces=audit_docx.NS,
            )
        )
        for text_node in caption.xpath(".//w:t", namespaces=audit_docx.NS):
            text_node.text = (text_node.text or "").replace(" [tab:main-symbols]", "")

    write_docx_mutation(missing_caption_tag, remove_caption_tag)
    missing_tag_report = audit_docx.audit(missing_caption_tag)
    missing_tag_check = next(
        item for item in missing_tag_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert missing_tag_check["ok"] is False
    assert missing_tag_check["caption_tags"] == []
    assert missing_tag_check["caption_tag_unique"] is False

    duplicate_caption_tag = tmp_path / "paper_duplicate_symbol_caption_tag.docx"

    def duplicate_caption_label(root):
        caption = next(
            paragraph for paragraph in root.xpath(".//w:body/w:p", namespaces=audit_docx.NS)
            if "[tab:main-symbols]" in audit_docx.element_text(paragraph)
            and paragraph.xpath(
                './w:pPr/w:pStyle[contains(translate(@w:val, "ABCDEFGHIJKLMNOPQRSTUVWXYZ", "abcdefghijklmnopqrstuvwxyz"), "huaweitablecaption")]',
                namespaces=audit_docx.NS,
            )
        )
        text_node = caption.xpath(".//w:t", namespaces=audit_docx.NS)[-1]
        text_node.text = (text_node.text or "") + " [tab:main-symbols]"

    write_docx_mutation(duplicate_caption_tag, duplicate_caption_label)
    duplicate_tag_report = audit_docx.audit(duplicate_caption_tag)
    duplicate_tag_check = next(
        item for item in duplicate_tag_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert duplicate_tag_check["ok"] is False
    assert duplicate_tag_check["caption_tags"] == ["tab:main-symbols", "tab:main-symbols"]
    assert duplicate_tag_check["caption_tag_unique"] is False

    native_ref = tmp_path / "paper_native_symbol_ref.docx"

    def replace_visible_tags_with_native_ref(root):
        caption = next(
            paragraph for paragraph in root.xpath(".//w:body/w:p", namespaces=audit_docx.NS)
            if "[tab:main-symbols]" in audit_docx.element_text(paragraph)
            and paragraph.xpath(
                './w:pPr/w:pStyle[contains(translate(@w:val, "ABCDEFGHIJKLMNOPQRSTUVWXYZ", "abcdefghijklmnopqrstuvwxyz"), "huaweitablecaption")]',
                namespaces=audit_docx.NS,
            )
        )
        for text_node in caption.xpath(".//w:t", namespaces=audit_docx.NS):
            text_node.text = (text_node.text or "").replace(" [tab:main-symbols]", "")
        bookmark_start = audit_docx.etree.Element(audit_docx.W + "bookmarkStart")
        bookmark_start.set(audit_docx.W + "id", "77")
        bookmark_start.set(audit_docx.W + "name", "MainSymbolCaption")
        bookmark_end = audit_docx.etree.Element(audit_docx.W + "bookmarkEnd")
        bookmark_end.set(audit_docx.W + "id", "77")
        caption.insert(1, bookmark_start)
        caption.append(bookmark_end)

        reference = next(
            paragraph for paragraph in root.xpath(".//w:body/w:p", namespaces=audit_docx.NS)
            if "全文主要符号见表" in audit_docx.element_text(paragraph)
        )
        text_nodes = reference.xpath(".//w:t", namespaces=audit_docx.NS)
        text_nodes[0].text = "全文主要符号见表所示。"
        for text_node in text_nodes[1:]:
            text_node.text = ""
        run = audit_docx.etree.SubElement(reference, audit_docx.W + "r")
        instruction = audit_docx.etree.SubElement(run, audit_docx.W + "instrText")
        instruction.text = " REF MainSymbolCaption \\h "

    write_docx_mutation(native_ref, replace_visible_tags_with_native_ref)
    native_report = audit_docx.audit(native_ref)
    assert native_report["status"] == "PASS", native_report["failures"]
    native_check = next(
        item for item in native_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert native_check["caption_tags"] == []
    assert native_check["caption_bookmarks"] == ["MainSymbolCaption"]
    assert native_check["references"][0]["ref_names"] == ["MainSymbolCaption"]

    fake_caption = tmp_path / "paper_fake_symbol_caption.docx"

    def replace_caption_with_misstyled_reference(root):
        body = root.find("w:body", audit_docx.NS)
        caption = next(
            paragraph for paragraph in body.findall("w:p", audit_docx.NS)
            if "符号" in audit_docx.element_text(paragraph)
            and paragraph.xpath(
                './w:pPr/w:pStyle[contains(translate(@w:val, "ABCDEFGHIJKLMNOPQRSTUVWXYZ", "abcdefghijklmnopqrstuvwxyz"), "huaweitablecaption")]',
                namespaces=audit_docx.NS,
            )
        )
        body.remove(caption)
        reference = next(
            paragraph for paragraph in body.findall("w:p", audit_docx.NS)
            if "全文主要符号见表" in audit_docx.element_text(paragraph)
        )
        properties = reference.find("w:pPr", audit_docx.NS)
        if properties is None:
            properties = audit_docx.etree.Element(audit_docx.W + "pPr")
            reference.insert(0, properties)
        style = properties.find("w:pStyle", audit_docx.NS)
        if style is None:
            style = audit_docx.etree.SubElement(properties, audit_docx.W + "pStyle")
        style.set(audit_docx.W + "val", "HuaweiTableCaption")

    write_docx_mutation(fake_caption, replace_caption_with_misstyled_reference)
    fake_caption_report = audit_docx.audit(fake_caption)
    fake_caption_check = next(
        item for item in fake_caption_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert fake_caption_check["ok"] is False
    assert fake_caption_check["caption_ok"] is False

    missing_reference = tmp_path / "paper_missing_symbol_reference.docx"

    def remove_symbol_reference(root):
        body = root.find("w:body", audit_docx.NS)
        reference = next(
            paragraph for paragraph in body.findall("w:p", audit_docx.NS)
            if "全文主要符号见表" in audit_docx.element_text(paragraph)
        )
        body.remove(reference)

    write_docx_mutation(missing_reference, remove_symbol_reference)
    missing_reference_report = audit_docx.audit(missing_reference)
    missing_reference_check = next(
        item for item in missing_reference_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert missing_reference_check["ok"] is False
    assert missing_reference_check["reference_before_table"] is False

    stale_number_reference = tmp_path / "paper_stale_symbol_table_number.docx"

    def replace_traceable_reference_with_plain_number(root):
        reference = next(
            paragraph for paragraph in root.xpath(".//w:body/w:p", namespaces=audit_docx.NS)
            if "全文主要符号见表" in audit_docx.element_text(paragraph)
        )
        text_nodes = reference.xpath(".//w:t", namespaces=audit_docx.NS)
        text_nodes[0].text = "全文主要符号见表 1。"
        for text_node in text_nodes[1:]:
            text_node.text = ""

    write_docx_mutation(stale_number_reference, replace_traceable_reference_with_plain_number)
    stale_reference_report = audit_docx.audit(stale_number_reference)
    stale_reference_check = next(
        item for item in stale_reference_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert stale_reference_check["ok"] is False
    assert stale_reference_check["reference_before_table"] is False

    missing_bookmark_reference = tmp_path / "paper_missing_symbol_bookmark.docx"

    def replace_traceable_reference_with_missing_bookmark(root):
        reference = next(
            paragraph for paragraph in root.xpath(".//w:body/w:p", namespaces=audit_docx.NS)
            if "全文主要符号见表" in audit_docx.element_text(paragraph)
        )
        text_nodes = reference.xpath(".//w:t", namespaces=audit_docx.NS)
        text_nodes[0].text = "全文主要符号见表所示。"
        for text_node in text_nodes[1:]:
            text_node.text = ""
        run = audit_docx.etree.SubElement(reference, audit_docx.W + "r")
        instruction = audit_docx.etree.SubElement(run, audit_docx.W + "instrText")
        instruction.text = " REF DefinitelyMissingBookmark \\h "

    write_docx_mutation(
        missing_bookmark_reference,
        replace_traceable_reference_with_missing_bookmark,
    )
    missing_bookmark_report = audit_docx.audit(missing_bookmark_reference)
    missing_bookmark_check = next(
        item for item in missing_bookmark_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert missing_bookmark_check["ok"] is False
    assert missing_bookmark_check["reference_before_table"] is False

    symbols_before_assumptions = tmp_path / "paper_symbols_before_assumptions.docx"

    def insert_assumptions_after_symbols(root):
        body = root.find("w:body", audit_docx.NS)
        children = list(body)
        problem_heading = next(
            paragraph for paragraph in children
            if paragraph.tag == audit_docx.W + "p"
            and "问题重述" in audit_docx.element_text(paragraph)
        )
        heading = deepcopy(problem_heading)
        text_nodes = heading.xpath(".//w:t", namespaces=audit_docx.NS)
        text_nodes[0].text = "2 模型假设"
        for text_node in text_nodes[1:]:
            text_node.text = ""
        body.insert(children.index(problem_heading), heading)

    write_docx_mutation(symbols_before_assumptions, insert_assumptions_after_symbols)
    order_report = audit_docx.audit(symbols_before_assumptions)
    order_check = next(
        item for item in order_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert order_check["ok"] is False
    assert order_check["after_existing_preliminary_sections"] is False

    incomplete_rows = tmp_path / "paper_incomplete_symbol_rows.docx"

    def clear_meaning_cells(root):
        table = root.xpath(".//w:tbl", namespaces=audit_docx.NS)[0]
        for row in table.xpath("./w:tr[position() > 1]", namespaces=audit_docx.NS):
            cells = row.xpath("./w:tc", namespaces=audit_docx.NS)
            for text_node in cells[1].xpath(".//w:t", namespaces=audit_docx.NS):
                text_node.text = ""

    write_docx_mutation(incomplete_rows, clear_meaning_cells)
    incomplete_report = audit_docx.audit(incomplete_rows)
    incomplete_check = next(
        item for item in incomplete_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert incomplete_check["ok"] is False
    assert incomplete_check["data_row_count"] == 0
    assert len(incomplete_check["incomplete_data_rows"]) == 2

    empty_units = tmp_path / "paper_empty_symbol_units.docx"

    def clear_unit_cells(root):
        table = root.xpath(".//w:tbl", namespaces=audit_docx.NS)[0]
        for row in table.xpath("./w:tr[position() > 1]", namespaces=audit_docx.NS):
            cells = row.xpath("./w:tc", namespaces=audit_docx.NS)
            for text_node in cells[2].xpath(".//w:t", namespaces=audit_docx.NS):
                text_node.text = ""

    write_docx_mutation(empty_units, clear_unit_cells)
    empty_units_report = audit_docx.audit(empty_units)
    empty_units_check = next(
        item for item in empty_units_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert empty_units_check["ok"] is False
    assert empty_units_check["expected_column_count"] == 3
    assert len(empty_units_check["incomplete_data_rows"]) == 2

    missing_unit_cell = tmp_path / "paper_missing_symbol_unit_cell.docx"

    def remove_unit_cell(root):
        table = root.xpath(".//w:tbl", namespaces=audit_docx.NS)[0]
        row = table.xpath("./w:tr[position() = 2]", namespaces=audit_docx.NS)[0]
        row.remove(row.xpath("./w:tc", namespaces=audit_docx.NS)[-1])

    write_docx_mutation(missing_unit_cell, remove_unit_cell)
    missing_unit_report = audit_docx.audit(missing_unit_cell)
    missing_unit_check = next(
        item for item in missing_unit_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert missing_unit_check["ok"] is False
    assert missing_unit_check["incomplete_data_rows"]

    placeholder_symbol = tmp_path / "paper_placeholder_symbol.docx"

    def inject_english_placeholder(root):
        table = root.xpath(".//w:tbl", namespaces=audit_docx.NS)[0]
        first_data_cell = table.xpath("./w:tr[position() = 2]/w:tc[1]", namespaces=audit_docx.NS)[0]
        text_nodes = first_data_cell.xpath(".//w:t", namespaces=audit_docx.NS)
        text_nodes[0].text = "REPLACE_WITH_ACTUAL_SYMBOL"
        for text_node in text_nodes[1:]:
            text_node.text = ""

    write_docx_mutation(placeholder_symbol, inject_english_placeholder)
    placeholder_report = audit_docx.audit(placeholder_symbol)
    placeholder_check = next(
        item for item in placeholder_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert placeholder_check["ok"] is False
    assert placeholder_check["placeholders"]

    single_column = tmp_path / "paper_single_column_symbol_table.docx"

    def merge_each_symbol_row_into_one_cell(root):
        table = root.xpath(".//w:tbl", namespaces=audit_docx.NS)[0]
        for row in table.xpath("./w:tr", namespaces=audit_docx.NS):
            cells = row.xpath("./w:tc", namespaces=audit_docx.NS)
            combined = " ".join(audit_docx.element_text(cell) for cell in cells)
            text_nodes = cells[0].xpath(".//w:t", namespaces=audit_docx.NS)
            text_nodes[0].text = combined
            for text_node in text_nodes[1:]:
                text_node.text = ""
            for cell in cells[1:]:
                row.remove(cell)

    write_docx_mutation(single_column, merge_each_symbol_row_into_one_cell)
    single_column_report = audit_docx.audit(single_column)
    single_column_check = next(
        item for item in single_column_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert single_column_check["ok"] is False
    assert single_column_check["expected_column_count"] == 1
    assert single_column_check["header_ok"] is False

    horizontal_grid = tmp_path / "paper_horizontal_grid_symbol_table.docx"

    def add_data_cell_rules(root):
        table = root.xpath(".//w:tbl", namespaces=audit_docx.NS)[0]
        for cell in table.xpath("./w:tr[position() > 1]/w:tc", namespaces=audit_docx.NS):
            properties = cell.find("w:tcPr", audit_docx.NS)
            borders = properties.find("w:tcBorders", audit_docx.NS)
            if borders is None:
                borders = audit_docx.etree.SubElement(properties, audit_docx.W + "tcBorders")
            bottom = audit_docx.etree.SubElement(borders, audit_docx.W + "bottom")
            bottom.set(audit_docx.W + "val", "single")

    write_docx_mutation(horizontal_grid, add_data_cell_rules)
    horizontal_report = audit_docx.audit(horizontal_grid)
    horizontal_check = next(
        item for item in horizontal_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert horizontal_check["ok"] is False
    assert horizontal_check["visible_data_cell_horizontal_borders"]

    duplicate_symbols = tmp_path / "paper_duplicate_symbol_sections.docx"

    def duplicate_symbol_section(root):
        body = root.find("w:body", audit_docx.NS)
        children = list(body)
        start = next(
            index for index, child in enumerate(children)
            if child.tag == audit_docx.W + "p"
            and audit_docx.SYMBOL_HEADING_RE.fullmatch(audit_docx.element_text(child))
        )
        level = audit_docx.paragraph_outline_level(children[start])
        end = next(
            index for index, child in enumerate(children[start + 1:], start + 1)
            if child.tag == audit_docx.W + "p"
            and audit_docx.paragraph_outline_level(child) is not None
            and audit_docx.paragraph_outline_level(child) <= level
        )
        for offset, child in enumerate(children[start:end]):
            body.insert(end + offset, deepcopy(child))

    write_docx_mutation(duplicate_symbols, duplicate_symbol_section)
    duplicate_report = audit_docx.audit(duplicate_symbols)
    duplicate_check = next(
        item for item in duplicate_report["checks"]
        if item["code"] == "main_symbol_glossary_is_complete_three_line_table"
    )
    assert duplicate_check["ok"] is False
    assert duplicate_check["symbol_heading_count"] == 2

    alternate_heading = tmp_path / "paper_alternate_symbol_heading.docx"

    def rename_symbol_heading(root):
        paragraph = next(
            item for item in root.xpath(".//w:body/w:p", namespaces=audit_docx.NS)
            if audit_docx.SYMBOL_HEADING_RE.fullmatch(audit_docx.element_text(item))
        )
        text_node = paragraph.xpath(".//w:t", namespaces=audit_docx.NS)[0]
        text_node.text = "1 变量与符号说明"

    write_docx_mutation(alternate_heading, rename_symbol_heading)
    alternate_report = audit_docx.audit(alternate_heading)
    assert alternate_report["status"] == "PASS", alternate_report["failures"]

    missing_heading = tmp_path / "paper_missing_appendix_heading.docx"
    original = "附录 A 可复现代码框架".encode("utf-8")
    replacement = "可复现代码框架".encode("utf-8")
    replaced = False
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(missing_heading, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                updated = data.replace(original, replacement)
                replaced = updated != data
                data = updated
            target.writestr(item, data)
    assert replaced
    missing_report = audit_docx.audit(missing_heading)
    terminal_check = next(
        item for item in missing_report["checks"]
        if item["code"] == "references_and_appendices_start_new_pages"
    )
    assert terminal_check["ok"] is True
    assert terminal_check["missing_roles"] == []
    assert terminal_check["optional_missing_roles"] == ["appendix"]


def test_copy_assets_validates_every_source_before_writing(tmp_path):
    build_latex = load_script("hardening_build_latex", SCRIPTS / "build_latex.py")
    template = tmp_path / "template"
    output = tmp_path / "output"
    (template / "figures").mkdir(parents=True)
    template.mkdir(exist_ok=True)
    for name in ("gmcmthesis.cls", "gmcm.bst", "reference.bib", "gmcm-title.sty"):
        (template / name).write_text(name, encoding="utf-8")
    with pytest.raises(FileNotFoundError, match="模板资产不完整"):
        build_latex.copy_assets(template, output)
    assert not output.exists() or not any(output.iterdir())


def test_latex_manifest_enforces_three_to_six_unique_keywords(tmp_path):
    build_latex = load_script("hardening_build_latex_keywords", SCRIPTS / "build_latex.py")
    schema = json.loads(
        (REPO_ROOT / "assets" / "paper-template" / "agent_manifest.schema.json").read_text(
            encoding="utf-8"
        )
    )
    keyword_schema = schema["properties"]["keywords"]
    assert keyword_schema["minItems"] == 3
    assert keyword_schema["maxItems"] == 6
    assert keyword_schema["uniqueItems"] is True

    manifest_path = tmp_path / "manifest.json"
    base = {
        "title": "模型研究",
        "abstract_tex_path": "abstract.tex",
        "appendix_pseudocode": {
            "required": False,
            "reason": "本测试仅检查关键词解析，不涉及程序求解或数值实验。",
        },
        "chapters": [
            {
                "chapter_id": "c1",
                "title": "问题重述",
                "role": "problem",
                "order": 1,
                "tex_path": "chapter.tex",
            }
        ],
    }

    def load(keywords):
        manifest_path.write_text(
            json.dumps({**base, "keywords": keywords}, ensure_ascii=False),
            encoding="utf-8",
        )
        return build_latex.load_manifest(manifest_path)

    assert load([" 建模 ", "验证", "优化"])["keywords"] == ["建模", "验证", "优化"]
    assert len(load(["一", "二", "三", "四", "五", "六"])["keywords"]) == 6
    for keywords, message in (
        (["一", "二"], "3–6"),
        (["一", "二", "三", "四", "五", "六", "七"], "3–6"),
        (["一", "二", "  "], "空字符串"),
        (["模型Ａ", "模型A", "验证"], "规范化后不得重复"),
        (["Model", "model", "验证"], "规范化后不得重复"),
        (["一", "二", 3], "字符串数组"),
    ):
        with pytest.raises(ValueError, match=message):
            load(keywords)


def test_appendix_pseudocode_policy_is_explicit_and_fail_closed(tmp_path):
    from jsonschema import Draft202012Validator

    build_latex = load_script("hardening_build_latex_appendix_policy", SCRIPTS / "build_latex.py")
    build_docx = load_script("hardening_build_docx_appendix_policy", SCRIPTS / "build_docx.py")
    schema = json.loads(
        (REPO_ROOT / "assets" / "paper-template" / "agent_manifest.schema.json").read_text(
            encoding="utf-8"
        )
    )
    validator = Draft202012Validator(schema)
    assert "appendix_pseudocode" in schema["required"]

    (tmp_path / "abstract.tex").write_text("摘要。", encoding="utf-8")
    (tmp_path / "chapter.tex").write_text("正文。", encoding="utf-8")
    (tmp_path / "references.tex").write_text("参考文献。", encoding="utf-8")
    (tmp_path / "appendix.tex").write_text("补充推导。", encoding="utf-8")
    manifest_path = tmp_path / "manifest.json"
    base = {
        "title": "解析模型研究",
        "keywords": ["解析", "推导", "验证"],
        "abstract_tex_path": "abstract.tex",
        "appendix_pseudocode": {
            "required": False,
            "reason": "本文仅作纯解析推导，不依赖程序求解或数值实验。",
        },
        "chapters": [
            {
                "chapter_id": "c1",
                "title": "问题分析",
                "role": "problem",
                "order": 1,
                "tex_path": "chapter.tex",
            },
            {
                "chapter_id": "refs",
                "title": "参考文献",
                "role": "references",
                "order": 2,
                "tex_path": "references.tex",
            }
        ],
    }

    def write(payload):
        manifest_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    assert not list(validator.iter_errors(base))
    write(base)
    assert build_latex.load_manifest(manifest_path)["appendix_pseudocode"]["required"] is False
    assert build_docx.read_input(manifest_path)["appendix_pseudocode"]["required"] is False

    missing = deepcopy(base)
    missing.pop("appendix_pseudocode")
    assert list(validator.iter_errors(missing))
    write(missing)
    for loader in (build_latex.load_manifest, build_docx.read_input):
        with pytest.raises(ValueError, match="appendix_pseudocode"):
            loader(manifest_path)

    vague = deepcopy(base)
    vague["appendix_pseudocode"]["reason"] = "这是一个占位理由而已"
    assert list(validator.iter_errors(vague))
    write(vague)
    for loader in (build_latex.load_manifest, build_docx.read_input):
        with pytest.raises(ValueError, match="充分具体"):
            loader(manifest_path)

    for sentinel_reason in (
        "REPLACE_WITH_ACTUAL：本文不涉及程序求解。",
        "TODO：本文仅作纯解析推导，不依赖数值程序。",
        "待替换：本文不涉及程序求解或数值实验。",
    ):
        sentinel = deepcopy(base)
        sentinel["appendix_pseudocode"]["reason"] = sentinel_reason
        assert list(validator.iter_errors(sentinel))
        write(sentinel)
        for loader in (build_latex.load_manifest, build_docx.read_input):
            with pytest.raises(ValueError, match="充分具体"):
                loader(manifest_path)

    required_without_appendix = deepcopy(base)
    required_without_appendix["appendix_pseudocode"] = {"required": True}
    assert list(validator.iter_errors(required_without_appendix))
    write(required_without_appendix)
    for loader in (build_latex.load_manifest, build_docx.read_input):
        with pytest.raises(ValueError, match="appendix 章节"):
            loader(manifest_path)

    required_with_appendix = deepcopy(required_without_appendix)
    required_with_appendix["chapters"].append({
        "chapter_id": "appendix",
        "title": "附录",
        "role": "appendix",
        "order": 3,
        "tex_path": "appendix.tex",
    })
    assert not list(validator.iter_errors(required_with_appendix))
    write(required_with_appendix)
    assert build_latex.load_manifest(manifest_path)["appendix_pseudocode"]["required"] is True
    assert build_docx.read_input(manifest_path)["appendix_pseudocode"]["required"] is True


def test_manifest_schema_requires_exactly_one_references_chapter():
    from jsonschema import Draft202012Validator

    schema = json.loads(
        (REPO_ROOT / "assets" / "paper-template" / "agent_manifest.schema.json").read_text(
            encoding="utf-8"
        )
    )
    Draft202012Validator.check_schema(schema)
    validator = Draft202012Validator(schema)
    manifest = {
        "title": "参考文献章节约束测试",
        "keywords": ["建模", "验证", "引用"],
        "abstract_tex_path": "abstract.tex",
        "appendix_pseudocode": {
            "required": False,
            "reason": "本文仅作纯解析推导，不依赖程序求解或数值实验。",
        },
        "chapters": [
            {
                "chapter_id": "body",
                "title": "问题分析",
                "role": "preliminary",
                "order": 1,
                "tex_path": "body.tex",
            },
            {
                "chapter_id": "refs",
                "title": "参考文献",
                "role": "references",
                "order": 2,
                "tex_path": "references.tex",
            },
        ],
    }
    assert not list(validator.iter_errors(manifest))

    no_references = deepcopy(manifest)
    no_references["chapters"] = no_references["chapters"][:-1]
    assert list(validator.iter_errors(no_references))

    duplicate_references = deepcopy(manifest)
    duplicate_references["chapters"].append({
        "chapter_id": "refs-duplicate",
        "title": "参考资料",
        "role": "references",
        "order": 3,
        "tex_path": "references-duplicate.tex",
    })
    assert list(validator.iter_errors(duplicate_references))


def test_latex_builder_output_passes_source_audit(tmp_path):
    build_latex = load_script("hardening_build_latex_e2e", SCRIPTS / "build_latex.py")
    audit_tex = load_script("hardening_audit_tex_e2e", SCRIPTS / "audit_tex.py")
    (tmp_path / "abstract.tex").write_text("摘要包含可复核的主要模型与结论。", encoding="utf-8")
    (tmp_path / "symbols.tex").write_text(
        "\n".join(
            [
                r"\section{主要符号说明}",
                r"全文主要符号见表\ref{tab:main-symbols}。",
                r"\begin{table}[htbp]",
                r"\caption{全文主要符号、定义与单位}",
                r"\label{tab:main-symbols}",
                r"\begin{tabular}{lll}",
                r"\toprule",
                r"符号 & 定义 & 单位 \\",
                r"\midrule",
                r"$x$ & 决策变量 & -- \\",
                r"$c$ & 单位成本 & 元 \\",
                r"\bottomrule",
                r"\end{tabular}",
                r"\end{table}",
            ]
        ),
        encoding="utf-8",
    )
    (tmp_path / "chapter.tex").write_text(
        r"\input{symbols.tex}" + "\n" + r"\section{问题重述}" + "\n" + "正文内容。",
        encoding="utf-8",
    )
    (tmp_path / "references.tex").write_text(
        "\n".join(
            [
                r"\begin{thebibliography}{9}",
                r"\bibitem{official} 官方格式规范。",
                r"\end{thebibliography}",
            ]
        ),
        encoding="utf-8",
    )
    valid_appendix = "\n".join(
        [
            r"\section{复现说明}",
            "对应脚本：solve.py；结果文件：results.csv。",
            r"\begin{Python}{可复现求解流程}",
            'config = read_json("config.json")',
            'data = read_csv("data.csv")',
            'train, validation, test = split_by_event(data, seed=config["seed"])',
            "model = fit(train, config)",
            "assert validate(model, validation)",
            "metrics = evaluate(model, test)",
            'write_csv(metrics, "results.csv")',
            r"\end{Python}",
        ]
    )
    (tmp_path / "appendix.tex").write_text(valid_appendix, encoding="utf-8")
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(
        json.dumps(
            {
                "title": "模型研究",
                "keywords": ["建模", "验证", "优化"],
                "abstract_tex_path": "abstract.tex",
                "appendix_pseudocode": {"required": True},
                "chapters": [
                    {"chapter_id": "c1", "title": "问题重述", "role": "problem", "order": 1, "tex_path": "chapter.tex"},
                    {"chapter_id": "refs", "title": "参考文献", "role": "references", "order": 2, "tex_path": "references.tex"},
                    {"chapter_id": "appendix", "title": "附录", "role": "appendix", "order": 3, "tex_path": "appendix.tex"},
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    config = {"contest": {"edition_cn": "二十三", "verified_against_official_rules": True}}
    manifest = build_latex.load_manifest(manifest_path)
    main_text, inputs = build_latex.build_main(manifest, tmp_path, config)
    assert inputs == ["abstract.tex", "chapter.tex", "references.tex", "appendix.tex"]
    assert r"\clearpage" + "\n" + r"\input{references.tex}" in main_text
    assert r"\clearpage" + "\n" + r"\appendix" + "\n" + r"\input{appendix.tex}" in main_text
    main_path = tmp_path / "paper.tex"
    main_path.write_text(main_text, encoding="utf-8")
    report = audit_tex.audit(manifest_path, main_path)
    assert report["status"] == "PASS", json.dumps(report["failures"], ensure_ascii=False, indent=2)

    layout_main = tmp_path / "paper_layout_override.tex"
    layout_main.write_text(
        main_text.replace(
            r"\pagestyle{plain}",
            r"\pagestyle{plain}" + "\n" + r"\newgeometry{margin=1cm}",
            1,
        ),
        encoding="utf-8",
    )
    layout_report = audit_tex.audit(manifest_path, layout_main)
    layout_check = next(
        item for item in layout_report["checks"]
        if item["code"] == "main_does_not_override_official_geometry"
    )
    assert layout_check["ok"] is False
    assert layout_check["overrides"]

    formal_blank_report = audit_tex.audit(
        manifest_path, main_path, require_filled_cover=True
    )
    assert "identity_cover_fields_filled_for_submission" in {
        item["code"] for item in formal_blank_report["failures"]
    }
    invisible_main_text = main_text
    for command in ("schoolname", "baominghao", "membera", "memberb", "memberc"):
        invisible_main_text = invisible_main_text.replace(
            rf"\{command}{{}}", rf"\{command}{{\quad}}"
        )
    invisible_main = tmp_path / "paper_invisible_cover.tex"
    invisible_main.write_text(invisible_main_text, encoding="utf-8")
    invisible_report = audit_tex.audit(
        manifest_path, invisible_main, require_filled_cover=True
    )
    assert "identity_cover_fields_filled_for_submission" in {
        item["code"] for item in invisible_report["failures"]
    }
    filled_main_text = main_text
    for command, value in {
        "schoolname": "甲大学",
        "baominghao": "TEAM-001",
        "membera": "甲",
        "memberb": "乙",
        "memberc": "丙",
    }.items():
        filled_main_text = filled_main_text.replace(
            rf"\{command}{{}}", rf"\{command}{{{value}}}"
        )
    filled_main = tmp_path / "paper_filled_cover.tex"
    filled_main.write_text(filled_main_text, encoding="utf-8")
    formal_filled_report = audit_tex.audit(
        manifest_path, filled_main, require_filled_cover=True
    )
    assert formal_filled_report["status"] == "PASS", json.dumps(
        formal_filled_report["failures"], ensure_ascii=False, indent=2
    )

    nested_code = valid_appendix.split("\n", 1)[1]
    (tmp_path / "appendix_code.tex").write_text(nested_code, encoding="utf-8")
    (tmp_path / "appendix.tex").write_text(
        r"\section{复现说明}" + "\n" + r"\input{appendix_code.tex}",
        encoding="utf-8",
    )
    nested_report = audit_tex.audit(manifest_path, main_path)
    assert nested_report["status"] == "PASS", json.dumps(
        nested_report["failures"], ensure_ascii=False, indent=2
    )

    (tmp_path / "appendix.tex").write_text(
        r"\section{复现说明}" + "\n" + "只有文字，没有代码式伪代码。",
        encoding="utf-8",
    )
    missing_code_report = audit_tex.audit(manifest_path, main_path)
    assert "appendix_has_code_style_pseudocode" in {
        item["code"] for item in missing_code_report["failures"]
    }

    (tmp_path / "appendix.tex").write_text(
        "\n".join(
            [
                r"\section{复现说明}",
                "对应脚本：solve.py；结果文件：results.csv。",
                r"\begin{Python}{空壳}",
                'print("hello")',
                r"\end{Python}",
            ]
        ),
        encoding="utf-8",
    )
    shallow_report = audit_tex.audit(manifest_path, main_path)
    assert next(
        item for item in shallow_report["checks"]
        if item["code"] == "appendix_has_code_style_pseudocode"
    )["ok"] is True
    assert "appendix_pseudocode_is_traceable" in {
        item["code"] for item in shallow_report["failures"]
    }

    generic_chain = "\n".join(
        [
            r"\section{复现说明}",
            "对应脚本：solve.py；结果文件：results.csv。",
            r"\begin{Python}{通用占位链}",
            'config = read_json("config.json")',
            'raw = read_csv("data.csv")',
            'data = audit_and_clean(raw)',
            'features = build_features(data)',
            'baseline = fit(features)',
            'result = solve_model(baseline)',
            'metrics = evaluate(result)',
            'write_csv(metrics, "results.csv")',
            r"\end{Python}",
        ]
    )
    (tmp_path / "appendix.tex").write_text(generic_chain, encoding="utf-8")
    generic_chain_report = audit_tex.audit(manifest_path, main_path)
    generic_check = next(
        item for item in generic_chain_report["checks"]
        if item["code"] == "appendix_pseudocode_is_traceable"
    )
    assert generic_check["ok"] is False
    assert generic_check["generic_placeholder_calls"] == [
        "audit_and_clean", "build_features", "solve_model"
    ]

    mapped_chain = generic_chain.replace(
        "对应脚本：solve.py；结果文件：results.csv。",
        (
            "函数映射：preprocess.py::audit_and_clean；"
            "features.py::build_features；solve.py::solve_model；"
            "结果文件：results.csv。"
        ),
    )
    (tmp_path / "appendix.tex").write_text(mapped_chain, encoding="utf-8")
    mapped_chain_report = audit_tex.audit(manifest_path, main_path)
    assert mapped_chain_report["status"] == "PASS", json.dumps(
        mapped_chain_report["failures"], ensure_ascii=False, indent=2
    )

    manifest_data = json.loads(manifest_path.read_text(encoding="utf-8"))
    original_manifest_data = deepcopy(manifest_data)
    manifest_data["appendix_pseudocode"] = {
        "required": False,
        "reason": "本例仅验证纯解析推导，不依赖程序求解或数值实验。",
    }
    manifest_path.write_text(json.dumps(manifest_data, ensure_ascii=False), encoding="utf-8")
    waiver_report = audit_tex.audit(manifest_path, main_path)
    assert waiver_report["status"] == "PASS", json.dumps(
        waiver_report["failures"], ensure_ascii=False, indent=2
    )

    (tmp_path / "appendix.tex").write_text(generic_chain, encoding="utf-8")
    malformed_waiver_report = audit_tex.audit(manifest_path, main_path)
    assert "appendix_pseudocode_is_traceable" in {
        item["code"] for item in malformed_waiver_report["failures"]
    }

    analytic_manifest = deepcopy(manifest_data)
    analytic_manifest["chapters"] = [
        chapter for chapter in analytic_manifest["chapters"]
        if chapter["role"] != "appendix"
    ]
    manifest_path.write_text(json.dumps(analytic_manifest, ensure_ascii=False), encoding="utf-8")
    analytic_main = tmp_path / "paper_analytic.tex"
    analytic_main.write_text(
        main_text.replace(
            r"\clearpage" + "\n" + r"\appendix" + "\n" + r"\input{appendix.tex}",
            "",
        ),
        encoding="utf-8",
    )
    analytic_report = audit_tex.audit(manifest_path, analytic_main)
    assert analytic_report["status"] == "PASS", json.dumps(
        analytic_report["failures"], ensure_ascii=False, indent=2
    )

    manifest_data = deepcopy(original_manifest_data)
    manifest_data["appendix_pseudocode"] = {
        "required": False,
        "reason": "太短",
    }
    (tmp_path / "appendix.tex").write_text(mapped_chain, encoding="utf-8")
    manifest_path.write_text(json.dumps(manifest_data, ensure_ascii=False), encoding="utf-8")
    bad_waiver_report = audit_tex.audit(manifest_path, main_path)
    assert "appendix_pseudocode_policy_valid" in {
        item["code"] for item in bad_waiver_report["failures"]
    }

    missing_policy = deepcopy(original_manifest_data)
    missing_policy.pop("appendix_pseudocode")
    manifest_path.write_text(json.dumps(missing_policy, ensure_ascii=False), encoding="utf-8")
    (tmp_path / "appendix.tex").write_text(valid_appendix, encoding="utf-8")
    missing_policy_report = audit_tex.audit(manifest_path, main_path)
    assert "appendix_pseudocode_policy_valid" in {
        item["code"] for item in missing_policy_report["failures"]
    }

    required_without_appendix = deepcopy(original_manifest_data)
    required_without_appendix["chapters"] = [
        chapter for chapter in required_without_appendix["chapters"]
        if chapter["role"] != "appendix"
    ]
    manifest_path.write_text(
        json.dumps(required_without_appendix, ensure_ascii=False), encoding="utf-8"
    )
    required_without_appendix_report = audit_tex.audit(manifest_path, analytic_main)
    assert "appendix_pseudocode_required_has_appendix" in {
        item["code"] for item in required_without_appendix_report["failures"]
    }

    manifest_path.write_text(
        json.dumps(original_manifest_data, ensure_ascii=False), encoding="utf-8"
    )

    bad_main = tmp_path / "paper_missing_reference_break.tex"
    bad_main.write_text(
        main_text.replace(r"\clearpage" + "\n" + r"\input{references.tex}", r"\input{references.tex}"),
        encoding="utf-8",
    )
    bad_report = audit_tex.audit(manifest_path, bad_main)
    assert "references_and_appendices_start_new_pages" in {
        item["code"] for item in bad_report["failures"]
    }


def test_main_symbol_glossary_contract_accepts_complete_table():
    audit_tex = load_script("hardening_symbol_glossary_valid", SCRIPTS / "audit_tex.py")
    source = r"""
\section{模型假设}
假设内容。
\subsection{主要符号说明}
全文主要符号见表\ref{tab:symbols}。
\begin{table}[htbp]
\caption{全文主要符号、定义与单位}
\label{tab:symbols}
\begin{tabularx}{\textwidth}{lXc}
\toprule
符号 & 定义 & 单位 \\
\midrule
$A\&B$ & 决策变量 & -- \\
$c$ & 单位成本 & 元 \\
\bottomrule
\end{tabularx}
\end{table}
\section{问题一：模型建立}
正文。
"""
    result = audit_tex.analyze_symbol_glossary(source)
    assert result["ok"] is True, result
    assert result["data_row_count"] == 2
    assert result["column_specs"] == ["lXc"]


def test_main_symbol_glossary_rejects_empty_or_missing_declared_fields():
    audit_tex = load_script("hardening_symbol_glossary_declared_fields", SCRIPTS / "audit_tex.py")
    source = r"""
\section{主要符号说明}
主要符号见表\ref{tab:symbols}。
\begin{table}
\caption{主要符号、定义与单位}\label{tab:symbols}
\begin{tabular}{lll}
\toprule
符号 & 定义 & 单位 \\
\midrule
$x$ & 决策变量 & -- \\
$c$ & 单位成本 & \\
\bottomrule
\end{tabular}
\end{table}
\section{问题一：模型建立}
正文。
"""
    result = audit_tex.analyze_symbol_glossary(source)
    assert result["ok"] is False
    assert result["expected_column_count"] == 3
    assert result["data_row_count"] == 1
    assert len(result["incomplete_data_rows"]) == 1
    assert any("每个已声明字段均不得留空" in item for item in result["failures"])


@pytest.mark.parametrize(
    ("case", "mutate", "expected_failure"),
    [
        (
            "missing",
            lambda text: r"\section{问题一：模型建立}" + "\n正文。",
            "未找到明确",
        ),
        (
            "commented-rules",
            lambda text: text.replace(r"\toprule", "% \\toprule")
            .replace(r"\midrule", "% \\midrule")
            .replace(r"\bottomrule", "% \\bottomrule"),
            "top/mid/bottomrule",
        ),
        (
            "grid",
            lambda text: text.replace("{ll}", "{|l|l|}").replace(r"\midrule", r"\midrule\hline"),
            "网格线",
        ),
        (
            "hidden-vrule-column",
            lambda text: text.replace("{ll}", r"{c!{\vrule}l}"),
            "网格线",
        ),
        (
            "tabular-star-optional-grid",
            lambda text: text.replace(
                r"\begin{tabular}{ll}",
                r"\begin{tabular*}{\textwidth}[t]{|l|l|}",
            ).replace(r"\end{tabular}", r"\end{tabular*}"),
            "列格式不得包含竖线",
        ),
        (
            "tabularx-optional-grid",
            lambda text: text.replace(
                r"\begin{tabular}{ll}",
                r"\begin{tabularx}{\textwidth}[t]{|l|X|}",
            ).replace(r"\end{tabular}", r"\end{tabularx}"),
            "列格式不得包含竖线",
        ),
        (
            "no-rows",
            lambda text: text.replace("$x$ & 决策变量 \\\\\n", "").replace(
                "$c$ & 单位成本 \\\\\n", ""
            ),
            "至少需要一条",
        ),
        (
            "empty-rows",
            lambda text: text.replace("$x$ & 决策变量", " ").replace("$c$ & 单位成本", " "),
            "至少需要一条",
        ),
        (
            "invisible-command-rows",
            lambda text: text.replace("$x$ & 决策变量", r"\text{} & \text{}").replace(
                "$c$ & 单位成本", r"\phantom{x} & \mbox{}"
            ),
            "至少需要一条",
        ),
        (
            "nested-phantom-rows",
            lambda text: text.replace(
                "$x$ & 决策变量", r"\phantom{\text{x}} & \hphantom{\mbox{含义}}"
            ).replace(
                "$c$ & 单位成本", r"\vphantom{\mathbf{c}} & \phantom{\text{成本}}"
            ),
            "至少需要一条",
        ),
        (
            "zero-width-wrapper-rows",
            lambda text: text.replace(
                "$x$ & 决策变量", r"\smash{\phantom{x}} & \rlap{}"
            ).replace(
                "$c$ & 单位成本", r"\null & \llap{}"
            ),
            "至少需要一条",
        ),
        (
            "spacing-only-rows",
            lambda text: text.replace(
                "$x$ & 决策变量", r"\hspace{1em} & \quad"
            ).replace("$c$ & 单位成本", r"\kern1em & \qquad"),
            "至少需要一条",
        ),
        (
            "symbol-spacing-only-rows",
            lambda text: text.replace("$x$ & 决策变量", r"~ & ~").replace(
                "$c$ & 单位成本", r"\, & \;"
            ),
            "至少需要一条",
        ),
        (
            "fill-only-rows",
            lambda text: text.replace("$x$ & 决策变量", r"\hfill & \vfill").replace(
                "$c$ & 单位成本", r"\relax & \null"
            ),
            "至少需要一条",
        ),
        (
            "zero-rule-only-rows",
            lambda text: text.replace(
                "$x$ & 决策变量", r"\rule{0pt}{1em} & \rule{1em}{0pt}"
            ).replace(
                "$c$ & 单位成本", r"\rule[1pt]{0.0pt}{2em} & \strut"
            ),
            "至少需要一条",
        ),
        (
            "partial-row-after-valid-rows",
            lambda text: text.replace(
                r"\bottomrule",
                r"$z$ & \\" + "\n" + r"\bottomrule",
            ),
            "每个已声明字段均不得留空",
        ),
        (
            "extra-rule-inside-data",
            lambda text: text.replace(
                "$x$ & 决策变量 \\\\\n",
                "$x$ & 决策变量 \\\\\n" + r"\midrule" + "\n",
            ),
            "真实数据区不得夹入额外",
        ),
        (
            "special-rule-inside-data",
            lambda text: text.replace(
                "$x$ & 决策变量 \\\\\n",
                "$x$ & 决策变量 \\\\\n" + r"\specialrule{1pt}{0pt}{0pt}" + "\n",
            ),
            "真实数据区不得夹入额外",
        ),
        (
            "special-rule-inside-header",
            lambda text: text.replace(
                "符号 & 含义 \\\\\n",
                "符号 & 含义 \\\\\n" + r"\specialrule{1pt}{0pt}{0pt}" + "\n",
            ),
            "不得使用 cmidrule/specialrule/hhline",
        ),
        (
            "cmidrule-inside-header",
            lambda text: text.replace(
                "符号 & 含义 \\\\\n",
                "符号 & 含义 \\\\\n" + r"\cmidrule{1-2}" + "\n",
            ),
            "不得使用 cmidrule/specialrule/hhline",
        ),
        (
            "special-rule-after-bottom",
            lambda text: text.replace(
                r"\end{tabular}",
                r"\specialrule{1pt}{0pt}{0pt}" + "\n" + r"\end{tabular}",
            ),
            "不得使用 cmidrule/specialrule/hhline",
        ),
        (
            "missing-inner-tabular",
            lambda text: text.replace(r"\begin{tabular}{ll}", "").replace(r"\end{tabular}", ""),
            "实际 tabular/tabularx",
        ),
        (
            "split-rules-across-tabulars",
            lambda text: text.replace(
                r"\midrule",
                r"\end{tabular}\begin{tabular}{ll}\midrule",
            ),
            "实际 tabular/tabularx",
        ),
        (
            "placeholder-row",
            lambda text: text.replace("$x$ & 决策变量", "待替换符号 & 待替换含义"),
            "占位文本",
        ),
        (
            "combined-or-empty-header-cells",
            lambda text: text.replace(
                "符号 & 含义", r"符号与含义 & \phantom{单位}"
            ),
            "两个不同表头单元格",
        ),
        (
            "multiple-tables",
            lambda text: text.replace(
                r"\section{问题一：模型建立}",
                (
                    r"\begin{table}\caption{补充符号表}\label{tab:sym-extra}"
                    r"\begin{tabular}{ll}\toprule 符号 & 含义 \\ \midrule "
                    r"$y$ & 响应变量 \\ $z$ & 状态变量 \\ \bottomrule"
                    r"\end{tabular}\end{table}"
                    "\n"
                    r"\section{问题一：模型建立}"
                ),
            ),
            "必须且只能包含一张主符号表",
        ),
        (
            "multiple-symbol-headings",
            lambda text: text + "\n" + text,
            "只能有一个独立的主要符号说明标题",
        ),
        (
            "comment-environment",
            lambda text: r"\begin{comment}" + "\n" + text + "\n" + r"\end{comment}",
            "未找到明确",
        ),
        (
            "iffalse-region",
            lambda text: r"\iffalse" + "\n" + text + "\n" + r"\fi",
            "未找到明确",
        ),
        (
            "nested-iffalse-region",
            lambda text: (
                r"\iffalse" + "\n" + r"\iffalse hidden \fi" + "\n"
                + text + "\n" + r"\fi"
            ),
            "未找到明确",
        ),
        (
            "iffalse-else-formula",
            lambda text: (
                r"\iffalse hidden\else\begin{flalign}x&=1\end{flalign}\fi"
                + "\n" + text
            ),
            "首个展示型模型公式之前",
        ),
        (
            "iftrue-fake-else-table",
            lambda text: r"\iftrue active without table\else" + text + r"\fi",
            "未找到明确",
        ),
        (
            "unknown-conditional",
            lambda text: r"\ifnum1=1 " + text + r"\fi",
            "无法静态判定",
        ),
        (
            "resizebox",
            lambda text: text.replace(
                r"\begin{tabular}{ll}",
                r"\resizebox{\textwidth}{!}{\begin{tabular}{ll}",
            ).replace(r"\end{tabular}", r"\end{tabular}}"),
            "不得使用 resizebox/scalebox",
        ),
        (
            "tiny-font",
            lambda text: text.replace(r"\begin{tabular}{ll}", r"\tiny\begin{tabular}{ll}"),
            "不得使用 resizebox/scalebox",
        ),
        (
            "small-font",
            lambda text: text.replace(r"\begin{tabular}{ll}", r"\small\begin{tabular}{ll}"),
            "不得使用 resizebox/scalebox",
        ),
        (
            "small-before-table",
            lambda text: text.replace(r"\begin{table}", r"\small\begin{table}"),
            "不得使用 resizebox/scalebox",
        ),
        (
            "persistent-small-before-heading",
            lambda text: r"\small" + "\n" + text,
            "不得使用 resizebox/scalebox",
        ),
        (
            "after-question",
            lambda text: text.replace(
                r"\section{问题一：模型建立}" + "\n正文。\n",
                "",
            ).replace(
                r"\section{主要符号说明}",
                r"\section{问题一：模型建立}" + "\n正文。\n" + r"\section{主要符号说明}",
            ),
            "位于第一道编号问题之前",
        ),
        (
            "formula-before-symbols",
            lambda text: r"\section{模型假设}\begin{equation}x=1\end{equation}" + "\n" + text,
            "首个展示型模型公式之前",
        ),
        (
            "formula-between-heading-and-table",
            lambda text: text.replace(
                r"主要符号见表\ref{tab:sym}。",
                r"主要符号见表\ref{tab:sym}。\begin{equation}x=1\end{equation}",
            ),
            "首个展示型模型公式之前",
        ),
        (
            "formula-inside-float-before-tabular",
            lambda text: text.replace(
                r"\caption{主要符号说明}\label{tab:sym}",
                r"\caption{主要符号说明}\label{tab:sym}"
                r"\begin{equation}x=1\end{equation}",
            ),
            "首个展示型模型公式之前",
        ),
        (
            "dollar-display-before-table",
            lambda text: text.replace(
                r"主要符号见表\ref{tab:sym}。",
                r"主要符号见表\ref{tab:sym}。$$x=1$$",
            ),
            "首个展示型模型公式之前",
        ),
        (
            "displaymath-before-table",
            lambda text: text.replace(
                r"主要符号见表\ref{tab:sym}。",
                r"主要符号见表\ref{tab:sym}。\begin{displaymath}x=1\end{displaymath}",
            ),
            "首个展示型模型公式之前",
        ),
        (
            "eqnarray-before-table",
            lambda text: text.replace(
                r"主要符号见表\ref{tab:sym}。",
                r"主要符号见表\ref{tab:sym}。\begin{eqnarray}x&=&1\end{eqnarray}",
            ),
            "首个展示型模型公式之前",
        ),
        (
            "alignat-before-table",
            lambda text: r"\begin{alignat}{2}x&=1\end{alignat}" + "\n" + text,
            "首个展示型模型公式之前",
        ),
        (
            "symbols-before-assumptions",
            lambda text: text.replace(
                r"\section{问题一：模型建立}",
                r"\section{模型假设}" + "\n假设内容。\n" + r"\section{问题一：模型建立}",
            ),
            "必须位于这些章节之后",
        ),
    ],
)
def test_main_symbol_glossary_contract_rejects_malformed_cases(case, mutate, expected_failure):
    audit_tex = load_script(f"hardening_symbol_glossary_{case}", SCRIPTS / "audit_tex.py")
    valid = (
        r"\section{主要符号说明}" + "\n"
        + r"主要符号见表\ref{tab:sym}。" + "\n"
        + r"\begin{table}" + "\n"
        + r"\caption{主要符号说明}\label{tab:sym}" + "\n"
        + r"\begin{tabular}{ll}" + "\n"
        + r"\toprule" + "\n"
        + "符号 & 含义 \\\\\n"
        + r"\midrule" + "\n"
        + "$x$ & 决策变量 \\\\\n"
        + "$c$ & 单位成本 \\\\\n"
        + r"\bottomrule" + "\n"
        + r"\end{tabular}" + "\n"
        + r"\end{table}" + "\n"
        + r"\section{问题一：模型建立}" + "\n正文。\n"
    )
    result = audit_tex.analyze_symbol_glossary(mutate(valid))
    assert result["ok"] is False, case
    assert any(expected_failure in item for item in result["failures"]), result
    if case == "multiple-tables":
        assert result["tables_in_symbol_section"] == 2
        assert result["single_table_in_symbol_section"] is False
    if case in {"resizebox", "tiny-font", "small-font", "small-before-table"}:
        assert result["forbidden_scaling_or_tiny_commands"]


def test_symbol_glossary_ignores_preamble_macros_and_local_symbol_headings():
    audit_tex = load_script("hardening_symbol_glossary_scope", SCRIPTS / "audit_tex.py")
    preamble = r"""
\documentclass{article}
\newcommand{\unusedfake}{\section{主要符号说明}\begin{equation}z=0\end{equation}}
\begin{document}
"""
    body = r"""
\section{主要符号说明}
主要符号见表\ref{tab:sym}。
\begin{table}
\caption{主要符号说明}\label{tab:sym}
\begin{tabular}{ll}
\toprule
符号 & 含义 \\
\midrule
$x$ & 决策变量 \\
\bottomrule
\end{tabular}
\end{table}
\section{问题一：模型建立}
正文。
\subsection{问题一局部符号说明}
式中参数只在本问首次出现处定义。
\end{document}
"""
    result = audit_tex.analyze_symbol_glossary(preamble + body)
    assert result["ok"] is True, result
    assert result["symbol_heading_count"] == 1


def test_symbol_glossary_audit_uses_nested_input_expansion(tmp_path):
    audit_tex = load_script("hardening_symbol_glossary_nested", SCRIPTS / "audit_tex.py")
    (tmp_path / "symbols.tex").write_text(
        "\n".join(
            [
                r"\section{符号说明}",
                r"主要符号见表\ref{tab:sym}。",
                r"\begin{longtable}{ll}",
                r"\caption{主要符号说明}\label{tab:sym} \\",
                r"\toprule",
                r"符号 & 含义 \\",
                r"\midrule",
                r"$x$ & 决策变量 \\",
                r"$y$ & 响应变量 \\",
                r"\bottomrule",
                r"\end{longtable}",
            ]
        ),
        encoding="utf-8",
    )
    main = tmp_path / "main.tex"
    main.write_text(
        r"\input{symbols.tex}" + "\n" + r"\section{问题一：求解}" + "\n正文。",
        encoding="utf-8",
    )
    expanded, errors = audit_tex.expand_tex_in_order(main, tmp_path)
    assert errors == []
    assert audit_tex.analyze_symbol_glossary(expanded)["ok"] is True


@pytest.mark.parametrize(
    "extra",
    [
        r"{\small 这是一段已在表前闭合的局部说明。}" + "\n",
        "",
    ],
)
def test_symbol_glossary_ignores_out_of_scope_small_text(extra):
    audit_tex = load_script("hardening_symbol_glossary_font_scope", SCRIPTS / "audit_tex.py")
    source = (
        r"\section{主要符号说明}" + "\n"
        + extra
        + r"主要符号见表\ref{tab:sym}。" + "\n"
        + r"\begin{table}" + "\n"
        + r"\caption{主要符号说明}\label{tab:sym}" + "\n"
        + r"\begin{tabular}{ll}\toprule" + "\n"
        + "符号 & 含义 \\\\\n"
        + r"\midrule" + "\n"
        + "$x$ & 决策变量 \\\\\n$c$ & 单位成本 \\\\\n"
        + r"\bottomrule\end{tabular}\end{table}" + "\n"
        + (r"{\small 这是一段表后的局部说明。}" + "\n" if not extra else "")
        + r"\section{问题一：模型建立}" + "\n正文。"
    )
    result = audit_tex.analyze_symbol_glossary(source)
    assert result["ok"] is True, result


def test_symbol_glossary_accepts_standard_longtable_head_and_foot():
    audit_tex = load_script("hardening_symbol_glossary_longtable", SCRIPTS / "audit_tex.py")
    source = r"""
\section{主要符号说明}
主要符号见表\ref{tab:sym-long}。
\begin{longtable}{lll}
\caption{主要符号说明}\label{tab:sym-long} \\
\toprule
符号 & 含义 & 单位 \\
\midrule
\endfirsthead
\toprule
符号 & 含义 & 单位 \\
\midrule
\endhead
\midrule
\multicolumn{3}{r}{续下页} \\
\bottomrule
\endfoot
\bottomrule
\endlastfoot
$x$ & 决策变量 & -- \\
$c$ & 单位成本 & 元 \\
\end{longtable}
\section{问题一：模型建立}
正文。
"""
    result = audit_tex.analyze_symbol_glossary(source)
    assert result["ok"] is True, result
    assert result["data_row_count"] == 2
    assert result["bound_data_environment"].lower() == "longtable"


def test_numbered_question_sections_require_progressive_openers(tmp_path):
    audit_tex = load_script("hardening_audit_tex_question_openers", SCRIPTS / "audit_tex.py")

    valid = r"""
\section{问题一：容量配置}
问题一承接题面给定的逐时负荷与设备上限，需要确定满足可靠性约束的储能容量。为避免直接使用复杂算法掩盖可行性，本问先用线性规划建立成本基线，再用蒙特卡洛场景检验极端负荷；最终给出容量、成本区间和失负荷风险，并把冻结容量传给问题二的调度模型。
\subsection{模型建立}
正文。
\section{问题二：独立校准}
问题二与问题一的容量决策相互独立，因为这里只使用观测样本校准传感器偏差。本问先估计分段误差，再用留出样本比较校准前后的绝对误差；最终输出校准曲线及置信区间，作为全文独立的质量控制结果。
\begin{equation}x=1\end{equation}
"""
    findings = audit_tex.audit_question_openers(valid)
    assert len(findings) == 2
    assert all(item["ok"] for item in findings), findings

    generic = r"""
\section{第 1 问：任务名称}
本问主要针对题目提出的问题开展相关内容分析，目标是完成问题研究。为此先分析数据，再采用合适的模型和算法完成求解；最后得到相应结果并给出结论，为后续研究提供依据。
\subsection{求解}
正文。
"""
    generic_finding = audit_tex.audit_question_openers(generic)[0]
    assert generic_finding["ok"] is False
    assert generic_finding["generic_only"] is True

    late = r"""
\section{问题3：预测}
\begin{equation}y=x\end{equation}
本问使用历史序列预测未来需求。先拟合基线再检验误差，最终给出预测区间并传给后续优化。
"""
    assert audit_tex.audit_question_openers(late)[0]["ok"] is False

    comments_only = r"""
\section{问题四：优化}
% 本问回溯输入，先建立模型，再给出结果。
\subsection{模型}
正文。
"""
    assert audit_tex.audit_question_openers(comments_only)[0]["ok"] is False
    assert audit_tex.audit_question_openers(r"\section{问题重述}\subsection{题意}正文。") == []
    assert audit_tex.audit_question_openers(
        r"\appendix\section{Q2 全部候选}\subsection{表格}正文。"
    ) == []

    (tmp_path / "opener.tex").write_text(
        "问题一使用题面给出的需求序列与容量上限，需要确定可行配置。"
        "为避免复杂算法掩盖可行性，本问先建立线性规划基线，再用场景扰动检查鲁棒性；"
        "最终给出容量和成本区间，"
        "并将冻结配置传给问题二。",
        encoding="utf-8",
    )
    (tmp_path / "question.tex").write_text(
        r"\section{问题一：配置}" + "\n" + r"\input{opener.tex}" + "\n"
        + r"\subsection{模型}" + "\n正文。",
        encoding="utf-8",
    )
    (tmp_path / "paper.tex").write_text(r"\input{question.tex}", encoding="utf-8")
    expanded, errors = audit_tex.expand_tex_in_order(tmp_path / "paper.tex", tmp_path)
    assert not errors
    assert audit_tex.audit_question_openers(expanded)[0]["ok"] is True


def test_manifest_driven_question_openers_reject_vacuum_alias_and_missing_question():
    audit_tex = load_script("hardening_manifest_question_openers", SCRIPTS / "audit_tex.py")
    manifest = {
        "chapters": [
            {
                "chapter_id": "q1",
                "title": "问题一：容量配置",
                "aliases": ["Q1"],
                "role": "problem",
            },
            {
                "chapter_id": "q2",
                "title": "极端校准",
                "aliases": ["第 2 问"],
                "role": "problem",
            },
            {
                "chapter_id": "q5",
                "title": "Question 5: 贡献归因",
                "aliases": ["贡献归因"],
                "role": "problem",
            },
        ]
    }

    vacuum = audit_tex.audit_manifest_question_openers(manifest, "正文没有任何问题标题。")
    assert vacuum["ok"] is False
    assert vacuum["expected_numbers"] == [1, 2, 5]
    assert vacuum["actual_numbers"] == []
    assert vacuum["missing_numbers"] == [1, 2, 5]

    opener = (
        "本问承接题面给定的逐时负荷与设备上限，需要确定满足可靠性约束的储能容量。"
        "为避免复杂算法掩盖可行性，本问先用线性规划建立成本基线，再用蒙特卡洛场景检验极端负荷；"
        "最终给出容量、成本区间和失负荷风险，并把冻结配置传给后续调度。"
    )
    missing_five = audit_tex.audit_manifest_question_openers(
        manifest,
        r"\section{Q1：容量配置}" + opener
        + r"\section{极端校准}" + opener,
    )
    assert missing_five["ok"] is False
    assert missing_five["actual_numbers"] == [1, 2]
    assert missing_five["missing_numbers"] == [5]

    complete = audit_tex.audit_manifest_question_openers(
        manifest,
        r"\section{问题一：容量配置}" + opener
        + r"\section{极端校准}" + opener
        + r"\section{贡献归因}" + opener,
    )
    assert complete["ok"] is True, complete
    assert complete["actual_numbers"] == [1, 2, 5]


def test_manifest_question_openers_reject_conflicting_aliases_and_duplicates():
    audit_tex = load_script("hardening_manifest_question_aliases", SCRIPTS / "audit_tex.py")
    conflicting = {
        "chapters": [
            {
                "chapter_id": "q1",
                "title": "问题一：配置",
                "aliases": ["Q2"],
                "role": "problem",
            }
        ]
    }
    result = audit_tex.audit_manifest_question_openers(conflicting, "")
    assert result["ok"] is False
    assert "多个编号" in result["registration_errors"][0]

    duplicate = {
        "chapters": [
            {"chapter_id": "q1a", "title": "问题一：配置", "role": "problem"},
            {"chapter_id": "q1b", "title": "Question 1: 校准", "role": "problem"},
        ]
    }
    result = audit_tex.audit_manifest_question_openers(duplicate, "")
    assert result["ok"] is False
    assert "重复登记" in result["registration_errors"][0]


def test_unnumbered_paper_can_declare_question_openers_not_applicable():
    audit_tex = load_script("hardening_unnumbered_question_waiver", SCRIPTS / "audit_tex.py")
    manifest = {
        "question_openers": {
            "required": False,
            "reason": "全文按非编号研究主题组织，不存在逐问章节。",
        },
        "chapters": [
            {"chapter_id": "recap", "title": "问题重述", "role": "problem"},
            {"chapter_id": "method", "title": "统一模型", "role": "evaluation"},
        ],
    }
    result = audit_tex.audit_manifest_question_openers(
        manifest,
        r"\section{问题重述}正文。\section{统一模型}正文。",
    )
    assert result["ok"] is True, result
    assert result["applicable"] is False
    assert result["expected_numbers"] == []

    undeclared = {"chapters": manifest["chapters"]}
    assert audit_tex.audit_manifest_question_openers(
        undeclared, r"\section{问题重述}正文。"
    )["ok"] is True

    contradiction = audit_tex.audit_manifest_question_openers(
        manifest,
        r"\section{Q1：未登记问题}" + (
            "本问使用逐时需求与容量上限，需要确定储能配置。"
            "为避免复杂算法掩盖约束，本问先建立线性规划，再用场景扰动检验；"
            "最终给出容量、成本与风险区间。"
        ),
    )
    assert contradiction["ok"] is False
    assert contradiction["unexpected_numbers"] == [1]


def test_tex_audit_fails_when_manifest_expected_question_heading_is_absent(tmp_path):
    audit_tex = load_script("hardening_question_vacuum_integration", SCRIPTS / "audit_tex.py")
    (tmp_path / "abstract.tex").write_text("摘要。", encoding="utf-8")
    (tmp_path / "question.tex").write_text(
        r"\section{统一求解}正文直接进入模型。", encoding="utf-8"
    )
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "title": "逐问导读门禁测试",
                "keywords": ["导读", "审计", "建模"],
                "abstract_tex_path": "abstract.tex",
                "chapters": [
                    {
                        "chapter_id": "q7",
                        "title": "问题七：统一求解",
                        "role": "problem",
                        "order": 1,
                        "tex_path": "question.tex",
                    }
                ],
                "appendix_pseudocode": {
                    "required": False,
                    "reason": "本测试只验证章节导读结构，不涉及程序求解。",
                },
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    main = tmp_path / "main.tex"
    main.write_text(
        r"\input{abstract.tex}" + "\n" + r"\input{question.tex}",
        encoding="utf-8",
    )
    report = audit_tex.audit(manifest, main)
    check = next(
        item for item in report["checks"]
        if item["code"] == "question_sections_open_with_recap_route_and_deliverable"
    )
    assert check["ok"] is False
    assert check["expected_numbers"] == [7]
    assert check["actual_numbers"] == []
    assert check["missing_numbers"] == [7]


def test_question_opener_manifest_schema_supports_explicit_not_applicable():
    schema = json.loads(
        (REPO_ROOT / "assets" / "paper-template" / "agent_manifest.schema.json").read_text(
            encoding="utf-8"
        )
    )
    policy = schema["properties"]["question_openers"]
    assert policy["required"] == ["required"]
    assert policy["properties"]["required"]["type"] == "boolean"
    assert policy["properties"]["reason"]["minLength"] == 8
    assert policy["allOf"][0]["then"]["required"] == ["reason"]


def test_builders_enforce_question_opener_manifest_contract(tmp_path):
    from jsonschema import Draft202012Validator

    build_latex = load_script("hardening_build_latex_question_policy", SCRIPTS / "build_latex.py")
    build_docx = load_script("hardening_build_docx_question_policy", SCRIPTS / "build_docx.py")
    schema = json.loads(
        (REPO_ROOT / "assets" / "paper-template" / "agent_manifest.schema.json").read_text(
            encoding="utf-8"
        )
    )
    validator = Draft202012Validator(schema)
    (tmp_path / "abstract.tex").write_text("摘要。", encoding="utf-8")
    (tmp_path / "chapter.tex").write_text("正文。", encoding="utf-8")
    (tmp_path / "references.tex").write_text("参考文献。", encoding="utf-8")
    manifest_path = tmp_path / "manifest.json"
    base = {
        "title": "逐问合同测试",
        "keywords": ["导读", "审计", "建模"],
        "abstract_tex_path": "abstract.tex",
        "appendix_pseudocode": {
            "required": False,
            "reason": "本测试不涉及程序求解或数值实验。",
        },
        "question_openers": {
            "required": False,
            "reason": "全文按非编号研究主题组织，不存在逐问章节。",
        },
        "chapters": [
            {
                "chapter_id": "recap",
                "title": "问题重述",
                "role": "problem",
                "order": 1,
                "tex_path": "chapter.tex",
            },
            {
                "chapter_id": "refs",
                "title": "参考文献",
                "role": "references",
                "order": 2,
                "tex_path": "references.tex",
            }
        ],
    }

    def write(payload):
        manifest_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    def assert_loads(payload):
        assert not list(validator.iter_errors(payload))
        write(payload)
        assert build_latex.load_manifest(manifest_path)["title"] == "逐问合同测试"
        assert build_docx.read_input(manifest_path)["title"] == "逐问合同测试"

    def assert_rejected(payload):
        write(payload)
        for loader in (build_latex.load_manifest, build_docx.read_input):
            with pytest.raises(ValueError, match="question_openers"):
                loader(manifest_path)

    assert_loads(base)

    numbered = deepcopy(base)
    numbered["question_openers"] = {"required": True}
    numbered["chapters"][0]["title"] = "问题七：统一求解"
    assert_loads(numbered)

    alias_numbered = deepcopy(numbered)
    alias_numbered["chapters"][0]["title"] = "统一求解"
    alias_numbered["chapters"][0]["aliases"] = ["Question 7"]
    assert_loads(alias_numbered)

    inferred = deepcopy(numbered)
    inferred.pop("question_openers")
    assert_loads(inferred)

    invalid_cases = []
    unknown = deepcopy(base)
    unknown["question_openers"]["extra"] = True
    invalid_cases.append(unknown)
    wrong_type = deepcopy(base)
    wrong_type["question_openers"]["required"] = 0
    invalid_cases.append(wrong_type)
    non_string_reason = deepcopy(base)
    non_string_reason["question_openers"]["reason"] = 12345678
    invalid_cases.append(non_string_reason)
    short_reason = deepcopy(base)
    short_reason["question_openers"]["reason"] = "不适用"
    invalid_cases.append(short_reason)
    false_with_number = deepcopy(numbered)
    false_with_number["question_openers"] = {
        "required": False,
        "reason": "全文没有编号问题，因此不需要逐问导读。",
    }
    invalid_cases.append(false_with_number)
    true_without_number = deepcopy(base)
    true_without_number["question_openers"] = {"required": True}
    invalid_cases.append(true_without_number)

    for payload in invalid_cases:
        assert list(validator.iter_errors(payload))
        assert_rejected(payload)

    conflicting_alias = deepcopy(numbered)
    conflicting_alias["chapters"][0]["aliases"] = ["Q8"]
    assert not list(validator.iter_errors(conflicting_alias))
    assert_rejected(conflicting_alias)


def test_top_level_figures_require_substantive_narrative_bridges():
    audit_tex = load_script("hardening_figure_bridges", SCRIPTS / "audit_tex.py")
    first = r"""
\begin{figure}
\caption{基线误差分布}\label{fig:baseline}
\end{figure}
"""
    second = r"""
\begin{figure}
\caption{校正后的分时效比较}\label{fig:calibrated}
\end{figure}
"""
    one_paragraph = (
        "图一的误差分布显示强回波区偏差集中在东侧边缘，说明平流外推对快速生消过程仍有系统遗漏。"
        "在此基础上，为进一步验证残差校正是否只改善局部纹理，下一图比较十个时效的 CSI 与 RMSE 变化。"
    )
    result = audit_tex.analyze_large_figure_bridges(first + one_paragraph + second)
    assert result["ok"] is True, result
    assert result["figure_count"] == 2
    assert result["bridges"][0]["substantive_paragraph_count"] == 1

    two_paragraphs = (
        "图一的误差分布显示强回波区偏差集中在东侧边缘，这一空间结构说明单纯平流难以覆盖局地快速生消。"
        "\n\n"
        "在此基础上，为检验残差校正能否稳定改善各个预报时效，下一图比较十个时效的 CSI 与 RMSE 变化。"
    )
    result = audit_tex.analyze_large_figure_bridges(first + two_paragraphs + second)
    assert result["ok"] is True, result
    assert result["bridges"][0]["substantive_paragraph_count"] == 2

    adjacent = audit_tex.analyze_large_figure_bridges(first + r"\clearpage" + second)
    assert adjacent["ok"] is False
    assert adjacent["bridges"][0]["substantive_paragraph_count"] == 0

    generic = audit_tex.analyze_large_figure_bridges(
        first + "如图所示，下面给出下一图进行进一步分析。" + second
    )
    assert generic["ok"] is False

    concise = audit_tex.analyze_large_figure_bridges(
        first
        + "图一显示误差随时效增长，说明长期预测逐步退化；为验证该趋势，下一图比较各时效的 CSI。"
        + second
    )
    assert concise["ok"] is True, concise

    vague_filler = audit_tex.analyze_large_figure_bridges(
        first
        + "图中结果显示模型情况存在一些值得关注的变化，需要结合实际情况开展深入讨论。"
        "在此基础上，下一图进一步分析模型结果，以便获得更加全面可靠的认识。"
        + second
    )
    assert vague_filler["ok"] is False

    single = audit_tex.analyze_large_figure_bridges(first)
    assert single["ok"] is True
    assert single["applicable"] is False


def test_figure_bridge_audit_counts_wrapper_calls_not_wrapper_definition():
    audit_tex = load_script("hardening_figure_bridge_wrappers", SCRIPTS / "audit_tex.py")
    source = r"""
\newcommand{\evidencefigure}[3]{%
  \begin{figure}\caption{#2}\label{#3}\end{figure}}
\evidencefigure{a.png}{基线误差分布}{fig:a}
图一显示极端样本的误差主要集中在短时强回波边缘，说明固定平流速度不能解释局部增强。

在此基础上，为验证校正是否在独立时效上保持一致，下一图比较各时效的阈值命中率和均方根误差。
\frameworkfigure{b.png}{分时效验证结果}{fig:b}
"""
    result = audit_tex.analyze_large_figure_bridges(source)
    assert result["ok"] is True, result
    assert result["figure_count"] == 2
    assert [item["label"] for item in result["figures"]] == ["fig:a", "fig:b"]


def test_figure_bridge_audit_handles_custom_wrappers_titles_and_code_literals():
    audit_tex = load_script("hardening_figure_bridge_edges", SCRIPTS / "audit_tex.py")
    custom = r"""
\documentclass{article}
\newcommand{\paperfigure}[2]{%
  \begin{figure}\caption{#1}\label{#2}\end{figure}}
\begin{document}
\paperfigure{第一张结果图}{fig:one}
图一显示误差主要集中在长时效强回波边缘，说明当前校正仍受位置偏差限制。

在此基础上，为验证这一限制是否跨事件稳定存在，下一图比较事件级 RMSE 与 CSI 差异。
\paperfigure{第二张结果图}{fig:two}
\end{document}
"""
    custom_result = audit_tex.analyze_large_figure_bridges(custom)
    assert custom_result["ok"] is True, custom_result
    assert custom_result["figure_count"] == 2
    assert [item["label"] for item in custom_result["figures"]] == ["fig:one", "fig:two"]

    heading_only = custom.replace(
        "图一显示误差主要集中在长时效强回波边缘，说明当前校正仍受位置偏差限制。\n\n"
        "在此基础上，为验证这一限制是否跨事件稳定存在，下一图比较事件级 RMSE 与 CSI 差异。",
        r"\paragraph{图一显示误差集中在长时效边缘；在此基础上，下一图比较事件级 RMSE 与 CSI 差异。}",
    )
    assert audit_tex.analyze_large_figure_bridges(heading_only)["ok"] is False

    literals = r"""
\documentclass{article}
\begin{document}
\begin{lstlisting}
\iffalse
\begin{figure}\caption{代码字面量}\label{fig:not-real}\end{figure}
\end{document}
\end{lstlisting}
\verb|\evidencefigure{fake.png}{代码字面量}{fig:also-not-real}|
\lstinline|\begin{figure}\caption{行内代码}\end{figure}|
\begin{verbatim*}\begin{figure}\caption{逐字代码}\end{figure}\end{verbatim*}
\begin{figure}\caption{唯一真实大图}\label{fig:real}\end{figure}
\end{document}
"""
    literal_result = audit_tex.analyze_large_figure_bridges(literals)
    assert literal_result["ok"] is True, literal_result
    assert literal_result["figure_count"] == 1
    assert literal_result["figures"][0]["label"] == "fig:real"

    first = r"\begin{figure}\caption{第一图}\label{fig:first}\end{figure}"
    second = r"\begin{figure}\caption{第二图}\label{fig:second}\end{figure}"
    hidden_bridge = (
        "图一显示误差集中在长时效边缘，说明位置偏差仍然突出；"
        "在此基础上，为验证跨事件稳定性，下一图比较事件级 RMSE 与 CSI。"
    )
    invisible_variants = [
        rf"\phantom{{{hidden_bridge}}}",
        rf"{{\color{{white}} {hidden_bridge}}}",
        rf"\makebox[0pt][l]{{{hidden_bridge}}}",
        rf"\color{{white}} {hidden_bridge} \color{{black}}",
        rf"\color[RGB]{{255,255,255}} {hidden_bridge} \color{{black}}",
        rf"\textcolor[RGB]{{255,255,255}}{{{hidden_bridge}}}",
        rf"{{\fontsize{{0pt}}{{0pt}}\selectfont {hidden_bridge}}}",
        rf"\scalebox{{0}}{{{hidden_bridge}}}",
        rf"\resizebox{{0pt}}{{!}}{{{hidden_bridge}}}",
        rf"\footnote{{{hidden_bridge}}}",
        rf"\marginpar{{{hidden_bridge}}}",
        rf"\todo{{{hidden_bridge}}}",
    ]
    for invisible in invisible_variants:
        hidden_result = audit_tex.analyze_large_figure_bridges(
            first + invisible + second
        )
        assert hidden_result["ok"] is False, (invisible, hidden_result)


@pytest.mark.parametrize(
    ("definition", "first_call", "second_call"),
    [
        (
            r"\newcommand{\paperfigure}[3][0.8]{\begin{figure}\caption{#2}\label{#3}\end{figure}}",
            r"\paperfigure{第一张结果图}{fig:one}",
            r"\paperfigure[0.7]{第二张结果图}{fig:two}",
        ),
        (
            r"\newcommand\paperfigure[2]{\begin{figure}\caption{#1}\label{#2}\end{figure}}",
            r"\paperfigure{第一张结果图}{fig:one}",
            r"\paperfigure{第二张结果图}{fig:two}",
        ),
        (
            r"\NewDocumentCommand{\paperfigure}{m m}{\begin{figure}\caption{#1}\label{#2}\end{figure}}",
            r"\paperfigure{第一张结果图}{fig:one}",
            r"\paperfigure{第二张结果图}{fig:two}",
        ),
        (
            r"\newcommand{\paperfigure}{\begin{figure}\caption{固定结果图}\label{fig:fixed}\end{figure}}",
            r"\paperfigure",
            r"\paperfigure",
        ),
    ],
)
def test_figure_bridge_audit_handles_common_wrapper_definition_forms(
    definition, first_call, second_call
):
    audit_tex = load_script("hardening_figure_wrapper_forms", SCRIPTS / "audit_tex.py")
    bridge = (
        "图一显示误差主要集中在长时效强回波边缘，说明位置偏差仍然突出；"
        "在此基础上，为验证该偏差是否跨事件稳定，下一图比较事件级 RMSE 与 CSI。"
    )
    source = (
        "\\documentclass{article}\n"
        + definition
        + "\n\\begin{document}\n"
        + first_call
        + "\n"
        + bridge
        + "\n"
        + second_call
        + "\n\\end{document}\n"
    )
    result = audit_tex.analyze_large_figure_bridges(source)
    assert result["ok"] is True, result
    assert result["figure_count"] == 2
    assert len(result["bridges"]) == 1


def test_figure_bridge_audit_respects_redefined_builtin_and_numbered_figures_only():
    audit_tex = load_script("hardening_figure_wrapper_override", SCRIPTS / "audit_tex.py")
    bridge = (
        "图一显示误差随时效上升，说明长时预测逐步退化；"
        "在此基础上，为验证阈值命中是否同步下降，下一图比较 CSI 与 FAR。"
    )
    source = (
        r"\documentclass{article}"
        r"\newcommand{\evidencefigure}[4][0.8]{"
        r"\begin{figure}\caption{#3}\label{#4}\end{figure}}"
        r"\begin{document}"
        r"\begin{figure}\includegraphics{decoration.png}\end{figure}"
        r"\begin{figure}\caption*{不编号装饰图}\end{figure}"
        r"\evidencefigure{a.png}{第一张结果图}{fig:one}"
        + bridge
        + r"\evidencefigure[0.7]{b.png}{第二张结果图}{fig:two}"
        r"\end{document}"
    )
    result = audit_tex.analyze_large_figure_bridges(source)
    assert result["ok"] is True, result
    assert result["figure_count"] == 2
    assert [item["label"] for item in result["figures"]] == ["fig:one", "fig:two"]

    caption_without_label = audit_tex.analyze_large_figure_bridges(
        r"\begin{figure}\caption{有独立图号但无标签}\end{figure}"
    )
    assert caption_without_label["figure_count"] == 1


def test_figure_graphics_resolve_wrapper_arguments_not_definition_placeholders():
    audit_tex = load_script("hardening_figure_graphic_paths", SCRIPTS / "audit_tex.py")
    source = r"""
\documentclass{article}
\newcommand{\evidencefigure}[3]{%
  \begin{figure}\includegraphics{figures/#1.pdf}\caption{#2}\label{#3}\end{figure}}
\begin{document}
\evidencefigure{actual}{中文结果图}{fig:actual}
\end{document}
"""
    assert audit_tex._figure_graphic_paths(source) == ["figures/actual.pdf"]
    tikz = r"""
\newcommand{\paperfigure}[2]{%
  \begin{figure}\begin{tikzpicture}\draw (0,0)--(1,1);\end{tikzpicture}
  \caption{#1}\label{#2}\end{figure}}
\paperfigure{纯 TikZ 图}{fig:tikz}
"""
    assert audit_tex._figure_graphic_paths(tikz) == []


def test_tex_audit_checks_existing_and_missing_delegated_wrapper_graphics(tmp_path):
    audit_tex = load_script("hardening_nested_wrapper_graphics", SCRIPTS / "audit_tex.py")
    (tmp_path / "abstract.tex").write_text("摘要。", encoding="utf-8")
    (tmp_path / "chapter.tex").write_text(
        r"图像见图\ref{fig:nested}。"
        r"\outerfigure{actual}{委托包装图}{fig:nested}",
        encoding="utf-8",
    )
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "title": "委托包装图审计",
                "keywords": ["图路径"],
                "abstract_tex_path": "abstract.tex",
                "chapters": [
                    {
                        "chapter_id": "c1",
                        "title": "正文",
                        "role": "problem",
                        "order": 1,
                        "tex_path": "chapter.tex",
                    }
                ],
                "appendix_pseudocode": {
                    "required": False,
                    "reason": "本路径审计测试不涉及程序求解",
                },
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    main = tmp_path / "paper.tex"
    main.write_text(
        "\n".join(
            [
                r"\documentclass{article}",
                r"\newcommand{\basefigure}[3]{%",
                r"  \begin{figure}\includegraphics{figures/#1.png}%",
                r"  \caption{#2}\label{#3}\end{figure}}",
                r"\newcommand{\outerfigure}[3]{\basefigure{nested/#1}{#2}{#3}}",
                r"\schoolname{}\baominghao{}\membera{}\memberb{}\memberc{}",
                r"\begin{document}\makeidentitycover\maketitle",
                r"\begin{abstract}\input{abstract.tex}\end{abstract}",
                r"\input{chapter.tex}",
                r"\end{document}",
            ]
        ),
        encoding="utf-8",
    )
    image = tmp_path / "figures" / "nested" / "actual.png"
    image.parent.mkdir(parents=True)
    image.write_bytes(b"test image placeholder")

    existing_report = audit_tex.audit(manifest, main)
    existing_check = next(
        item for item in existing_report["checks"]
        if item["code"] == "graphics_paths_exist"
    )
    assert existing_check["ok"] is True
    expanded, expansion_errors = audit_tex.expand_tex_in_order(
        main, tmp_path, base=tmp_path
    )
    assert expansion_errors == []
    assert audit_tex._figure_graphic_paths(expanded) == [
        "figures/nested/actual.png"
    ]

    image.unlink()
    missing_report = audit_tex.audit(manifest, main)
    missing_check = next(
        item for item in missing_report["checks"]
        if item["code"] == "graphics_paths_exist"
    )
    assert missing_check["ok"] is False
    assert missing_check["missing"] == ["figures/nested/actual.png"]


def test_figure_bridge_counts_nested_math_captions():
    audit_tex = load_script("hardening_nested_figure_captions", SCRIPTS / "audit_tex.py")
    source = (
        r"\begin{figure}\caption{模型 $Z_{DR}$ 的比较}\label{fig:zdr}\end{figure}"
        r"\begin{figure}\caption[短题注]{模型 $K_{DP}$ 的比较}\label{fig:kdp}\end{figure}"
    )
    result = audit_tex.analyze_large_figure_bridges(source)
    assert result["figure_count"] == 2
    assert result["ok"] is False


def test_ab_comparison_requires_complete_consistent_measurements():
    ab_test = load_script("hardening_ab_test", FIGURE_SCRIPTS / "ab_test.py")
    with pytest.raises(ValueError, match="missing scenario"):
        ab_test.validate_result_set({}, "baseline")

    payload = {}
    for scenario in ab_test.SCENARIOS:
        total = len(scenario["expected"]["checks"])
        payload[scenario["id"]] = {"passed": total, "total": total, "pass_rate": 1.0}
    assert ab_test.validate_result_set(payload, "baseline") == payload
    payload[ab_test.SCENARIOS[0]["id"]]["pass_rate"] = 0.5
    with pytest.raises(ValueError, match="inconsistent"):
        ab_test.validate_result_set(payload, "baseline")


def test_generated_source_contract_distinguishes_pass_and_fail():
    runner = load_script("hardening_e2e_runner", FIGURE_SCRIPTS / "e2e_runner.py")
    scenario = next(item for item in runner.SCENARIOS if item.id == "S4_unknown_chart")
    bad = runner.score_script("print('not a figure')", scenario)
    assert bad["meets_threshold"] is False
    good_source = """
# cross-type inherit from a network asset
CATEGORICAL = ["#2166AC"]
FONT = "Arial"
import matplotlib.pyplot as plt
plt.plot([1, 2], [2, 3])
"""
    good = runner.score_script(good_source, scenario)
    assert good["meets_threshold"] is True


def test_tex_audit_handles_escaped_percent_and_rejects_nested_escape(tmp_path):
    audit_tex = load_script("hardening_audit_tex", SCRIPTS / "audit_tex.py")
    assert audit_tex.remove_comments(r"保留 50\% 内容 % 删除注释") == r"保留 50\% 内容 "

    abstract = tmp_path / "abstract.tex"
    chapter = tmp_path / "chapter.tex"
    secret = tmp_path.parent / "secret.tex"
    abstract.write_text("摘要", encoding="utf-8")
    chapter.write_text(r"\section{正文}" + "\n" + r"\input{../secret}", encoding="utf-8")
    secret.write_text("secret", encoding="utf-8")
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "title": "测试",
                "keywords": ["测试"],
                "abstract_tex_path": "abstract.tex",
                "appendix_pseudocode": {
                    "required": False,
                    "reason": "本测试只验证路径边界，不涉及程序求解或数值实验。",
                },
                "chapters": [
                    {"chapter_id": "c1", "title": "正文", "role": "problem", "order": 1, "tex_path": "chapter.tex"}
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    main = tmp_path / "paper.tex"
    main.write_text(
        "\n".join(
            [
                r"\schoolname{}\baominghao{}\membera{}\memberb{}\memberc{}",
                r"\begin{document}\makeidentitycover\maketitle",
                r"\begin{abstract}\input{abstract.tex}\end{abstract}",
                r"\input{chapter.tex}",
                r"\end{document}",
            ]
        ),
        encoding="utf-8",
    )
    report = audit_tex.audit(manifest, main)
    assert report["status"] == "FAIL"
    nested = next(item for item in report["checks"] if item["code"] == "nested_inputs_are_static_and_local")
    assert nested["ok"] is False


def test_tex_audit_rejects_shell_escape_in_main(tmp_path):
    audit_tex = load_script("hardening_audit_tex_shell", SCRIPTS / "audit_tex.py")
    abstract = tmp_path / "abstract.tex"
    chapter = tmp_path / "chapter.tex"
    abstract.write_text("摘要", encoding="utf-8")
    chapter.write_text(r"\section{正文}" + "\n正文", encoding="utf-8")
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "title": "测试",
                "keywords": ["测试"],
                "abstract_tex_path": "abstract.tex",
                "appendix_pseudocode": {
                    "required": False,
                    "reason": "本测试只验证命令安全，不涉及程序求解或数值实验。",
                },
                "chapters": [
                    {"chapter_id": "c1", "title": "正文", "role": "problem", "order": 1, "tex_path": "chapter.tex"}
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    main = tmp_path / "paper.tex"
    main.write_text(
        "\n".join(
            [
                r"\schoolname{}\baominghao{}\membera{}\memberb{}\memberc{}",
                r"\begin{document}\makeidentitycover\maketitle",
                r"\immediate\write18{whoami}",
                r"\begin{abstract}\input{abstract.tex}\end{abstract}",
                r"\input{chapter.tex}",
                r"\end{document}",
            ]
        ),
        encoding="utf-8",
    )
    report = audit_tex.audit(manifest, main)
    dangerous = next(
        item for item in report["checks"]
        if item["code"] == "no_dangerous_tex_file_or_shell_commands"
    )
    assert report["status"] == "FAIL"
    assert dangerous["ok"] is False
    assert any("write18" in command for command in dangerous["commands"])


def test_eval_runner_unknown_type_fails_without_traceback():
    process = subprocess.run(
        [
            sys.executable,
            str(FIGURE_SCRIPTS / "eval_runner.py"),
            "--type",
            "DefinitelyMissingFigureType",
        ],
        cwd=REPO_ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=30,
    )
    assert process.returncode == 1
    assert "not found" in process.stdout.lower()
    assert "traceback" not in (process.stdout + process.stderr).lower()


def test_submission_audit_refuses_unreadable_anonymous_input(tmp_path):
    empty = tmp_path / "empty.txt"
    empty.write_text("", encoding="utf-8")
    process = subprocess.run(
        [
            sys.executable,
            "-X",
            "utf8",
            str(SCRIPTS / "submission_audit.py"),
            "--paper",
            str(empty),
            "--cover-policy",
            "forbidden",
            "--ai-used",
            "none",
            "--json",
        ],
        cwd=REPO_ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=30,
    )
    assert process.returncode == 1
    report = json.loads(process.stdout)
    assert any(not item["ok"] and "文本可读取" in item["desc"] for item in report["checks"])


def test_submission_audit_preserves_escaped_percent_but_removes_comments():
    audit = load_script("hardening_submission_comments", SCRIPTS / "submission_audit.py")
    assert audit.strip_tex_comments(r"保留 95\% 结果 % 删除注释") == r"保留 95\% 结果 "
    assert audit.strip_tex_comments(r"换行\\% 删除注释") == "换行\\\\"


def test_static_figure_audit_is_complete_without_empirical_claim():
    runner = load_script("hardening_static_ab", FIGURE_SCRIPTS / "run_ab_tests.py")
    report = runner.run_all()
    assert report["summary"]["all_passed"] is True
    assert report["empirical_quality_claim"] is False
