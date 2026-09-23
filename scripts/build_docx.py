#!/usr/bin/env python3
# [来源] 移植自 v2.1 华为杯_论文规范模板/tools/，相对路径已改为 CLI 参数驱动，Wave3 验证编译链路。
"""Generate an optional derived DOCX from the TeX-first paper manifest.

The formal source of truth is the LaTeX ``论文.tex`` input chain. This
utility reads the same ``.tex`` fragments and makes a deliberately limited
plain-text/figure/table rendering for users who also need a Word derivative.
It never accepts Markdown chapter sources or inline prose fields, and it
never edits the official template in place.
"""
from __future__ import annotations

import argparse
from io import BytesIO
import json
import re
import time
import unicodedata
import zipfile
from pathlib import Path
from urllib.parse import unquote

from audit_tex import (
    MANIFEST_REASON_SENTINEL_RE,
    _balanced_macro_calls,
    question_opener_manifest_contract,
)



def load_contest_config(start: Path, explicit: Path | None = None) -> dict:
    if explicit is not None:
        candidate = explicit.resolve()
        if not candidate.is_file():
            raise FileNotFoundError(f"比赛配置不存在: {candidate}")
        return json.loads(candidate.read_text(encoding="utf-8-sig"))
    for base in [start.resolve(), *start.resolve().parents]:
        candidate = base / "比赛配置.json"
        if candidate.is_file():
            return json.loads(candidate.read_text(encoding="utf-8-sig"))
    raise FileNotFoundError("未找到根目录 比赛配置.json；请先运行 比赛当天初始化.bat 或指定 --contest-config")


def contest_edition(config: dict) -> str:
    contest = config.get("contest", {})
    if contest.get("verified_against_official_rules") is not True:
        raise ValueError("比赛配置中的届次尚未通过当届官方规则核对：请先更新 contest.edition_cn / edition_arabic，并将 verified_against_official_rules 设为 true")
    edition = str(contest.get("edition_cn", "")).strip()
    if not edition:
        raise ValueError("比赛配置.json 缺少 contest.edition_cn")
    return edition

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{" + W_NS + "}"
ROLES = {"preliminary", "problem", "evaluation", "conclusion", "references", "appendix"}

IDENTITY_RE = re.compile(
    r"(?:学校(?:名称)?|所属学校|学院(?:名称)?|实验室(?:名称)?|参赛队号|队号|"
    r"队员(?:姓名)?(?:\s*[123一二三])?|指导(?:教师|老师)|学号|姓名|(?:电子)?邮箱)\s*[:：]|"
    r"\b(?:school|student|team|member|advisor|e-?mail)\s*(?:name|id|number)?\s*[:：]|"
    r"C:\\Users\\",
    re.IGNORECASE,
)


def set_run_font(run, east_asia: str, size_pt: float, *, bold: bool = False, italic: bool = False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:hint"), "eastAsia")


def set_style_font(style, east_asia: str, size_pt: float, *, bold: bool = False, italic: bool = False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    # Do not inherit character emphasis from the official template.  The
    # competition body text is upright; inline emphasis is opt-in below.
    style.font.italic = italic
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:hint"), "eastAsia")


def set_spacing(paragraph, *, before=0, after=0, first_line=True):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.first_line_indent = Inches(1 / 3) if first_line else Inches(0)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def clear_element_content(element):
    """Remove runs/drawings/content while retaining paragraph properties."""
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)


def clear_body_keep_final_section(document):
    body = document._element.body
    final_sect = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not final_sect:
            body.remove(child)
    if final_sect is None:
        final_sect = OxmlElement("w:sectPr")
        body.append(final_sect)
    return document.sections[0]


def configure_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(1.75)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)
    section.different_first_page_header_footer = False
    section.start_type = 0  # WD_SECTION_START.CONTINUOUS; one section remains.
    sect_pr = section._sectPr
    page_numbering = sect_pr.find(qn("w:pgNumType"))
    if page_numbering is None:
        page_numbering = OxmlElement("w:pgNumType")
        sect_pr.append(page_numbering)
    page_numbering.set(qn("w:start"), "1")


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, "宋体", 10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_header_footer(section):
    header = section.header
    for p in header.paragraphs:
        clear_element_content(p._p)
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)

    footer = section.footer
    for p in footer.paragraphs:
        clear_element_content(p._p)
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_field(fp)


def add_style(document, name, east_asia, size, *, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, outline_level=None):
    styles = document.styles
    try:
        style = styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(style, east_asia, size, bold=bold)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    style.paragraph_format.first_line_indent = Inches(1 / 3) if align == WD_ALIGN_PARAGRAPH.JUSTIFY else Inches(0)
    style.paragraph_format.alignment = align
    # Avoid inheriting the template's automatic numbering when we write our own.
    ppr = style._element.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is not None:
        ppr.remove(num_pr)
    outline = ppr.find(qn("w:outlineLvl"))
    if outline is not None:
        ppr.remove(outline)
    if outline_level is not None:
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(outline_level))
        ppr.append(outline)
    return style


def configure_styles(document):
    normal = document.styles["Normal"]
    set_style_font(normal, "宋体", 12)
    normal.paragraph_format.first_line_indent = Inches(1 / 3)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    add_style(document, "Huawei Body", "宋体", 12)
    add_style(document, "Huawei Heading 1", "黑体", 14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, outline_level=0)
    add_style(document, "Huawei Heading 2", "宋体", 12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, outline_level=1)
    add_style(document, "Huawei Heading 3", "宋体", 12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, outline_level=2)
    add_style(document, "Huawei Appendix Code", "Times New Roman", 9, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_style(document, "Huawei Table Caption", "宋体", 12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_label_run(paragraph, text):
    run = paragraph.add_run(text)
    set_run_font(run, "隶书", 18)
    return run


def add_body_run(paragraph, text, *, bold=False, italic=False):
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", 12, bold=bold, italic=italic)
    return run


def add_body_paragraph(document, text, *, style="Huawei Body", first_line=True, space_before=0, space_after=0):
    p = document.add_paragraph(style=style)
    set_spacing(p, before=space_before, after=space_after, first_line=first_line)
    # TeX is converted to plain text before it reaches Word. Always write
    # upright runs here; LaTeX remains the authoritative formatting route.
    add_body_run(p, text.strip())
    return p


def add_heading(document, text, level, number=None, *, role=None):
    if number is not None:
        rendered = f"{number} {text}"
    else:
        rendered = text
    style = "Huawei Heading 1" if level == 1 else "Huawei Heading 2" if level == 2 else "Huawei Heading 3"
    p = document.add_paragraph(style=style)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(6 if level == 1 else 3)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 3)
    # The official format requires single line spacing for all Chinese text.
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.keep_with_next = True
    ppr = p._p.get_or_add_pPr()
    outline = ppr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        ppr.append(outline)
    outline.set(qn("w:val"), str(max(0, min(2, level - 1))))
    run = p.add_run(rendered)
    set_run_font(run, "黑体" if level == 1 else "宋体", 14 if level == 1 else 12, bold=level <= 2)
    p._p.set(qn("w:rsidR"), "00000000")
    return p, rendered


def enable_update_fields(document):
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _set_border(parent, edge, value, *, size=None):
    borders_tag = "w:tblBorders" if parent.tag == qn("w:tblPr") else "w:tcBorders"
    borders = parent.find(qn(borders_tag))
    if borders is None:
        borders = OxmlElement(borders_tag)
        parent.append(borders)
    border = borders.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        borders.append(border)
    border.set(qn("w:val"), value)
    if size is not None:
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")


def _apply_three_line_borders(table):
    table.style = None
    properties = table._tbl.tblPr
    for edge in ("left", "right", "insideH", "insideV"):
        _set_border(properties, edge, "nil")
    for edge in ("top", "bottom"):
        _set_border(properties, edge, "single", size=8)
    for cell in table.rows[0].cells:
        cell_properties = cell._tc.get_or_add_tcPr()
        _set_border(cell_properties, "bottom", "single", size=8)


def add_table(document, rows, caption="", labels=()):
    if not rows:
        return
    if caption:
        cap = document.add_paragraph(style="Huawei Table Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.first_line_indent = Inches(0)
        caption_text = tex_to_word_text(caption)
        tags = " ".join(f"[{label}]" for label in labels)
        add_body_run(cap, f"{caption_text} {tags}".strip(), bold=True)
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col in range(columns):
            cell = cells[col]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = values[col] if col < len(values) else ""
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for run in p.runs:
                    set_run_font(run, "宋体", 12, bold=row_index == 0)
    _apply_three_line_borders(table)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


TEX_IMAGE_RE = re.compile(r"\\includegraphics(?:\[[^]]*\])?\s*\{([^{}]+)\}")
TEX_CAPTION_RE = re.compile(r"\\caption(?:\[[^]]*\])?\s*\{([^{}]*)\}")
TEX_TABLE_LABEL_RE = re.compile(r"\\label\s*\{(tab:[^{}]+)\}", re.I)
TEX_HEADING_RE = re.compile(r"\\(section|subsection|subsubsection)\*?\s*\{([^{}]*)\}")
TEX_BEGIN_RE = re.compile(r"\\begin\s*\{([^{}]+)\}")
WORD_IMAGE_SUFFIXES = {".bmp", ".gif", ".jpeg", ".jpg", ".png", ".tif", ".tiff"}
WORD_FIGURE_WRAPPERS = ("evidencefigure", "frameworkfigure", "roadmapfigure")


def _replace_braced_command(text: str, command: str, replacement=None) -> str:
    """Replace simple one-level ``\\command{...}`` occurrences."""
    pattern = re.compile(rf"\\{command}\s*\{{([^{{}}]*)\}}")
    return pattern.sub(r"\1" if replacement is None else replacement, text)


def tex_to_word_text(text: str) -> str:
    """Convert common inline TeX to readable upright Word text.

    This is intentionally not a second TeX engine. Complex formulas remain
    readable as text in the optional derivative; the authoritative PDF is
    always produced by XeLaTeX from the same fragments.
    """
    text = text.replace("\\textbackslash{}", "\\")
    text = text.replace("\\textasciitilde{}", "~").replace("\\textasciicircum{}", "^")
    text = re.sub(r"\\frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}", r"(\1)/(\2)", text)
    text = re.sub(r"\\sqrt\s*\{([^{}]*)\}", r"sqrt(\1)", text)
    for command in ("texttt", "textbf", "textit", "emph", "textrm", "textsf", "mathrm", "mathbf", "mathit", "text"):
        text = _replace_braced_command(text, command)
    text = re.sub(r"\\(?:left|right|!|,|;|:|quad|qquad|enspace|hspace|vspace)\s*(?:\{[^{}]*\})?", " ", text)
    replacements = {
        r"\\times": "×", r"\\cdot": "·", r"\\leq": "≤", r"\\geq": "≥",
        r"\\neq": "≠", r"\\approx": "≈", r"\\in": "∈", r"\\notin": "∉",
        r"\\sum": "Σ", r"\\prod": "Π", r"\\sqrt": "√", r"\\infty": "∞",
        r"\\alpha": "α", r"\\beta": "β", r"\\gamma": "γ", r"\\delta": "δ",
        r"\\lambda": "λ", r"\\mu": "μ", r"\\sigma": "σ", r"\\pi": "π",
        r"\\rho": "ρ", r"\\tau": "τ", r"\\phi": "φ", r"\\omega": "ω",
        r"\\pm": "±", r"\\cdots": "…", r"\\ldots": "…",
    }
    for raw, value in replacements.items():
        text = re.sub(raw + r"(?![A-Za-z@])", value, text)
    text = re.sub(r"\\(?:label|ref|pageref|cite|eqref)\s*\{([^{}]*)\}", r"[\1]", text)
    text = re.sub(r"\\(?:small|footnotesize|normalsize|noindent|centering|par|protect|mbox|raisebox)(?:\s*\{[^{}]*\})?", "", text)
    text = text.replace("\\\\", " ")
    text = text.replace("\\_", "_").replace("\\%", "%").replace("\\&", "&")
    text = text.replace("\\#", "#").replace("\\{", "{").replace("\\}", "}")
    text = text.replace("~", " ")
    text = re.sub(r"\$+|\\\(|\\\)|\\\[|\\\]", "", text)
    text = re.sub(r"\\[A-Za-z@]+", "", text)
    text = text.replace("{", "").replace("}", "")
    return re.sub(r"\s+", " ", text).strip()


def _skip_tex_group(text: str, cursor: int, opening: str, closing: str) -> int | None:
    while cursor < len(text) and text[cursor].isspace():
        cursor += 1
    if cursor >= len(text) or text[cursor] != opening:
        return None
    depth = 1
    cursor += 1
    while cursor < len(text) and depth:
        if text[cursor] == opening and text[cursor - 1] != "\\":
            depth += 1
        elif text[cursor] == closing and text[cursor - 1] != "\\":
            depth -= 1
        cursor += 1
    return cursor if depth == 0 else None


def _extract_tabular_body(text: str) -> str:
    begin = re.search(
        r"\\begin\s*\{(?P<env>tabular\*?|tabularx|longtable)\}",
        text,
        re.IGNORECASE,
    )
    if begin is None:
        return text
    env = begin.group("env")
    cursor = begin.end()
    if env.lower() in {"tabular*", "tabularx"}:
        cursor = _skip_tex_group(text, cursor, "{", "}")
        if cursor is None:
            return text
    optional_end = _skip_tex_group(text, cursor, "[", "]")
    if optional_end is not None:
        cursor = optional_end
    cursor = _skip_tex_group(text, cursor, "{", "}")
    if cursor is None:
        return text
    end = re.search(
        rf"\\end\s*\{{{re.escape(env)}\}}",
        text[cursor:],
        re.IGNORECASE,
    )
    return text[cursor:cursor + end.start()] if end is not None else text[cursor:]


def tex_table_rows(text: str) -> list[list[str]]:
    """Extract a simple tabular/longtable body into Word rows."""
    text = _extract_tabular_body(text)
    if re.search(r"\\endfirsthead\b", text, re.IGNORECASE):
        first_head, continuation = re.split(
            r"\\endfirsthead\b", text, maxsplit=1, flags=re.IGNORECASE
        )
        marker_matches = list(re.finditer(
            r"\\end(?:head|foot|lastfoot)\b", continuation, re.IGNORECASE
        ))
        data = continuation[marker_matches[-1].end():] if marker_matches else continuation
        text = first_head + "\n" + data
    text = re.sub(r"(?m)^\\(?:centering|small|footnotesize|scriptsize)\s*$", "", text)
    text = re.sub(r"(?m)^\\caption\s*\{.*\}\s*$", "", text)
    text = re.sub(r"(?m)^\\label\s*\{.*\}\s*$", "", text)
    text = re.sub(r"\\(?:toprule|midrule|bottomrule|hline|endhead|endfirsthead|endfoot|endlastfoot)\b", "", text)
    rows: list[list[str]] = []
    for raw_row in re.split(r"\\\\", text):
        row = raw_row.strip()
        if not row or row.startswith("\\caption") or row.startswith("\\label"):
            continue
        if re.search(r"\\multicolumn\b", row, re.IGNORECASE):
            continue
        cells = [tex_to_word_text(cell) for cell in re.split(r"(?<!\\)&", row)]
        if any(cell != "" for cell in cells):
            rows.append(cells)
    return rows


def resolve_figure_path(base_dir: Path, raw_path: str) -> Path:
    """Resolve a TeX image path without allowing project-root escape."""
    clean_path = raw_path.strip()
    for escaped, literal in ((r"\ ", " "), (r"\#", "#"), (r"\%", "%"), (r"\&", "&")):
        clean_path = clean_path.replace(escaped, literal)
    requested = Path(clean_path)
    if requested.suffix and requested.suffix.lower() not in WORD_IMAGE_SUFFIXES | {".pdf"}:
        raise ValueError(
            f"Word 衍生稿不支持图片格式 {requested.suffix}: {raw_path}"
        )
    roots = (base_dir.resolve(), (base_dir / "figures").resolve())
    suffixes = ("",) if requested.suffix else (
        ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".pdf"
    )
    candidates: list[Path] = []
    for search_root in roots:
        for suffix in suffixes:
            candidate = (
                (search_root / requested)
                if requested.suffix else (search_root / requested).with_suffix(suffix)
            ).resolve()
            try:
                candidate.relative_to(base_dir.resolve())
            except ValueError as exc:
                raise ValueError(f"图片必须位于论文工作目录内: {raw_path}") from exc
            if candidate not in candidates:
                candidates.append(candidate)
    match = next((candidate for candidate in candidates if candidate.is_file()), None)
    if match is None:
        raise FileNotFoundError(f"图片不存在: {raw_path}")
    return match


def add_tex_figure(document, raw_path: str, caption: str, base_dir: Path):
    image_path = resolve_figure_path(base_dir, raw_path)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    p.paragraph_format.line_spacing = Pt(1)
    if image_path.suffix.lower() == ".pdf":
        import pymupdf as fitz

        with fitz.open(image_path) as pdf:
            if pdf.page_count < 1:
                raise ValueError(f"PDF 图片没有可渲染页面: {image_path}")
            pixmap = pdf[0].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image_source = BytesIO(pixmap.tobytes("png"))
        p.add_run().add_picture(image_source, width=Cm(15.5))
    else:
        p.add_run().add_picture(str(image_path), width=Cm(15.5))
    cap = document.add_paragraph(style="Huawei Table Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Inches(0)
    add_body_run(cap, tex_to_word_text(caption), bold=True)


def _normalize_production_figure_calls(content: str) -> str:
    """Put balanced built-in Figure wrapper calls on standalone logical lines."""
    calls = _balanced_macro_calls(content, WORD_FIGURE_WRAPPERS, 3)
    for call in reversed(calls):
        arguments = [re.sub(r"\s+", " ", value).strip() for value in call["args"]]
        replacement = (
            "\n\\" + call["name"]
            + "".join("{" + value + "}" for value in arguments)
            + "\n"
        )
        content = content[:call["start"]] + replacement + content[call["end"]:]
    return content


def add_tex_content(document, content: str, base_dir: Path, *, appendix=False, skip_first_section=True):
    """Render a TeX fragment as an optional, non-authoritative Word derivative."""
    content = _normalize_production_figure_calls(content)
    lines = content.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    paragraph_buffer: list[str] = []
    i = 0
    skipped_section = False

    def flush():
        if paragraph_buffer:
            text = tex_to_word_text(" ".join(paragraph_buffer))
            if text:
                add_body_paragraph(document, text, style="Huawei Appendix Code" if appendix else "Huawei Body", first_line=not appendix)
            paragraph_buffer.clear()

    while i < len(lines):
        raw = lines[i].strip()
        if not raw or raw.startswith("%"):
            flush()
            i += 1
            continue
        heading = TEX_HEADING_RE.fullmatch(raw)
        if heading:
            flush()
            kind, title = heading.groups()
            if kind == "section" and skip_first_section and not skipped_section:
                skipped_section = True
            else:
                level = {"section": 1, "subsection": 2, "subsubsection": 3}[kind]
                add_heading(document, tex_to_word_text(title), level)
            i += 1
            continue
        begin = TEX_BEGIN_RE.match(raw)
        if begin and begin.group(1) in {
            "table", "table*", "tabular", "tabular*", "tabularx", "longtable",
            "figure", "figure*",
        }:
            flush()
            env = begin.group(1)
            base_env = env.rstrip("*")
            block = []
            depth = 1
            i += 1
            while i < len(lines):
                if re.match(rf"^\\begin\s*\{{{re.escape(env)}\}}", lines[i].strip()):
                    depth += 1
                if re.fullmatch(rf"\\end\s*\{{{re.escape(env)}\}}", lines[i].strip()):
                    depth -= 1
                    if depth == 0:
                        break
                block.append(lines[i])
                i += 1
            block_text = "\n".join(block)
            if base_env == "figure":
                image = TEX_IMAGE_RE.search(block_text)
                caption = TEX_CAPTION_RE.search(block_text)
                if image:
                    add_tex_figure(document, image.group(1), caption.group(1) if caption else "", base_dir)
            else:
                caption = TEX_CAPTION_RE.search(block_text)
                add_table(
                    document,
                    tex_table_rows(block_text),
                    caption.group(1) if caption else "",
                    TEX_TABLE_LABEL_RE.findall(block_text),
                )
            i += 1
            continue
        if raw.startswith("\\begin{") or raw.startswith("\\end{"):
            if raw.startswith("\\begin{itemize") or raw.startswith("\\begin{enumerate"):
                flush()
                i += 1
                continue
            if raw.startswith("\\end{itemize") or raw.startswith("\\end{enumerate"):
                flush()
                i += 1
                continue
            if raw.startswith("\\begin{equation") or raw.startswith("\\begin{align") or raw.startswith("\\begin{gather"):
                flush()
                i += 1
                formula = []
                while i < len(lines) and not lines[i].lstrip().startswith("\\end{"):
                    formula.append(lines[i].strip())
                    i += 1
                if formula:
                    add_body_paragraph(document, tex_to_word_text(" ".join(formula)), first_line=False)
                i += 1
                continue
            i += 1
            continue
        if raw.startswith("\\item"):
            flush()
            paragraph = tex_to_word_text(re.sub(r"^\\item\s*", "", raw))
            if paragraph:
                add_body_paragraph(document, "• " + paragraph, first_line=False)
            i += 1
            continue
        wrapper_calls = _balanced_macro_calls(raw, WORD_FIGURE_WRAPPERS, 3)
        if wrapper_calls:
            cursor = 0
            for call in wrapper_calls:
                prefix = raw[cursor:call["start"]].strip()
                if prefix:
                    paragraph_buffer.append(prefix)
                    flush()
                add_tex_figure(
                    document,
                    call["args"][0],
                    call["args"][1],
                    base_dir,
                )
                cursor = call["end"]
            suffix = raw[cursor:].strip()
            if suffix:
                paragraph_buffer.append(suffix)
            i += 1
            continue
        image = TEX_IMAGE_RE.search(raw)
        if image:
            flush()
            caption = TEX_CAPTION_RE.search(raw)
            add_tex_figure(document, image.group(1), caption.group(1) if caption else "", base_dir)
            i += 1
            continue
        caption = TEX_CAPTION_RE.fullmatch(raw)
        if caption:
            flush()
            p = document.add_paragraph(style="Huawei Table Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            add_body_run(p, tex_to_word_text(caption.group(1)), bold=True)
            i += 1
            continue
        paragraph_buffer.append(raw)
        i += 1
    flush()


def resolve_tex_path(root: Path, raw: str, label: str) -> Path:
    path = (root / raw).resolve()
    try:
        path.relative_to(root.resolve())
    except ValueError as exc:
        raise ValueError(f"{label} 必须位于论文工作目录内: {raw}") from exc
    if path.suffix.lower() != ".tex":
        raise ValueError(f"{label} 必须使用 .tex，Markdown 不是论文源: {raw}")
    if not path.is_file():
        raise FileNotFoundError(f"{label} 不存在: {path}")
    return path


def read_tex_source(path: Path, root: Path, seen=None) -> str:
    """Read a fragment and expand active, static local ``\\input``/``\\include`` calls."""
    seen = set() if seen is None else seen
    path = path.resolve()
    if path in seen:
        raise ValueError(f"TeX 输入循环: {path}")
    seen.add(path)
    text = path.read_text(encoding="utf-8-sig")

    def expand(match):
        raw = match.group(1).strip()
        candidate = raw if raw.lower().endswith(".tex") else raw + ".tex"
        child = resolve_tex_path(root, candidate, f"TeX input in {path.name}")
        return read_tex_source(child, root, seen.copy())

    input_re = re.compile(r"(?<!\\)\\(?:input|include)\s*\{([^{}]+)\}")

    def expand_active_part(line: str) -> str:
        comment_at = None
        for index, character in enumerate(line):
            if character != "%":
                continue
            backslashes = 0
            cursor = index - 1
            while cursor >= 0 and line[cursor] == "\\":
                backslashes += 1
                cursor -= 1
            if backslashes % 2 == 0:
                comment_at = index
                break
        active = line if comment_at is None else line[:comment_at]
        comment = "" if comment_at is None else line[comment_at:]
        active = input_re.sub(expand, active)
        if re.search(r"(?<!\\)\\(?:input|include)\b", active):
            raise ValueError(f"存在无法静态展开的 TeX input/include: {path}")
        return active + comment

    expanded = "".join(expand_active_part(line) for line in text.splitlines(keepends=True))
    return expanded


def read_input(path: Path):
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(data, dict):
        raise ValueError("论文输入必须是 JSON 对象")
    for key in ("title", "abstract_tex_path", "keywords", "appendix_pseudocode", "chapters"):
        if key not in data:
            raise ValueError(f"输入缺少字段: {key}")
    forbidden = {"abstract", "content", "content_file"} & set(data)
    if forbidden:
        raise ValueError(f"TeX-first 输入禁止旧字段: {', '.join(sorted(forbidden))}")
    if not str(data["title"]).strip():
        raise ValueError("论文题目不能为空")
    if not isinstance(data["keywords"], list) or not all(
        isinstance(item, str) for item in data["keywords"]
    ):
        raise ValueError("关键词必须是字符串数组")
    keywords = [item.strip() for item in data["keywords"]]
    if any(not item for item in keywords):
        raise ValueError("关键词不能包含空字符串")
    if not 3 <= len(keywords) <= 6:
        raise ValueError("关键词必须提供 3–6 个")
    normalized_keywords = [
        re.sub(r"\s+", "", unicodedata.normalize("NFKC", item)).casefold()
        for item in keywords
    ]
    if len(set(normalized_keywords)) != len(normalized_keywords):
        raise ValueError("关键词规范化后不得重复")
    data["keywords"] = keywords
    root = path.parent.resolve()
    abstract_path = resolve_tex_path(root, str(data["abstract_tex_path"]), "abstract_tex_path")
    if not read_tex_source(abstract_path, root).strip():
        raise ValueError("摘要 TeX 内容不能为空")
    if not isinstance(data["chapters"], list) or not data["chapters"]:
        raise ValueError("chapters 必须是非空列表")
    identity_values = [data.get("title", ""), data.get("abstract_tex_path", "")]
    identity_values.extend(str(k) for k in data["keywords"])
    chapter_ids: set[str] = set()
    chapter_orders: set[int] = set()
    for chapter in data["chapters"]:
        if not isinstance(chapter, dict):
            raise ValueError("chapters 中的每一项都必须是对象")
        required = {"chapter_id", "title", "role", "order", "tex_path"}
        missing = sorted(required - set(chapter))
        if missing:
            raise ValueError(f"章节缺少字段: {', '.join(missing)}")
        if {"content", "content_file"} & set(chapter):
            raise ValueError(f"章节 {chapter.get('chapter_id', '?')} 禁止旧 content/content_file 字段")
        chapter_id = str(chapter["chapter_id"]).strip()
        if not chapter_id or chapter_id in chapter_ids:
            raise ValueError(f"chapter_id 不能为空且不得重复: {chapter_id!r}")
        chapter_ids.add(chapter_id)
        try:
            chapter_order = int(chapter["order"])
        except (TypeError, ValueError) as exc:
            raise ValueError(f"章节 {chapter_id} 的 order 必须是整数") from exc
        if chapter_order in chapter_orders:
            raise ValueError(f"章节 order 不得重复: {chapter_order}")
        chapter_orders.add(chapter_order)
        if not str(chapter["title"]).strip():
            raise ValueError(f"章节 {chapter_id} 的 title 不能为空")
        if chapter["role"] not in ROLES:
            raise ValueError(f"章节 {chapter_id} 的 role 无效: {chapter['role']}")
        try:
            level = int(chapter.get("level", 1))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"章节 {chapter_id} 的 level 必须是 1、2 或 3") from exc
        if level not in {1, 2, 3}:
            raise ValueError(f"章节 {chapter_id} 的 level 必须是 1、2 或 3")
        chapter_path = resolve_tex_path(root, str(chapter["tex_path"]), f"章节 {chapter['chapter_id']}")
        identity_values.extend([str(chapter.get("title", "")), str(chapter["tex_path"])])
        # Read now so the Word derivative fails early on malformed encoding or
        # a recursive input chain instead of producing a partial document.
        read_tex_source(chapter_path, root)
    orders = [int(chapter["order"]) for chapter in data["chapters"]]
    if orders != sorted(orders):
        raise ValueError("章节 order 必须严格递增")
    question_contract = question_opener_manifest_contract(data)
    if not question_contract["ok"]:
        raise ValueError("question_openers 配置无效: " + "; ".join(question_contract["errors"]))
    policy = data["appendix_pseudocode"]
    if not isinstance(policy, dict) or set(policy) - {"required", "reason"}:
        raise ValueError("appendix_pseudocode 必须是仅含 required/reason 的对象")
    if type(policy.get("required")) is not bool:
        raise ValueError("appendix_pseudocode.required 必须显式为布尔值")
    if "reason" in policy and not isinstance(policy["reason"], str):
        raise ValueError("appendix_pseudocode.reason 必须是字符串")
    if policy["required"] is False:
        reason = str(policy.get("reason", "")).strip()
        if len(reason) < 8 or MANIFEST_REASON_SENTINEL_RE.search(reason) or not re.search(
            r"纯解析|解析推导|理论推导|闭式|证明|不依赖|不涉及|无需|未使用|没有使用|"
            r"analytic|closed[- ]form|proof|without|does not",
            reason,
            re.IGNORECASE,
        ):
            raise ValueError("appendix_pseudocode.required=false 必须给出充分具体的不适用理由")
    if policy["required"] is True and not any(
        chapter["role"] == "appendix" for chapter in data["chapters"]
    ):
        raise ValueError("appendix_pseudocode.required=true 时必须声明 appendix 章节")
    leaked = [value for value in identity_values if IDENTITY_RE.search(value)]
    if leaked:
        raise ValueError("输入疑似含身份信息，请清理学校/队号/姓名/路径后再生成")
    return data


def remove_numbering_from_paragraph(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    num = ppr.find(qn("w:numPr"))
    if num is not None:
        ppr.remove(num)


def is_external_local_file_target(target: str, target_mode: str = "") -> bool:
    """Identify external relationships that expose a local/network path."""
    decoded = unquote(target).strip().replace("\\", "/")
    lower = decoded.lower()
    if lower.startswith("file:") or re.match(r"^[a-z]:/", lower):
        return True
    if decoded.startswith("//"):
        return True
    return target_mode.lower() == "external" and decoded.startswith("/")


def scrub_package(path: Path):
    """Remove template-only hidden payloads and personal core properties."""
    temp = path.with_suffix(path.suffix + ".cleaning")
    with zipfile.ZipFile(path, "r") as src, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            name = info.filename
            if (
                name.startswith("customXml/")
                or name.startswith("word/embeddings/")
                or name == "docProps/custom.xml"
            ):
                continue
            data = src.read(name)
            if name == "docProps/core.xml":
                root = etree_from_bytes(data)
                for tag in ("creator", "lastModifiedBy", "description", "subject", "keywords", "category"):
                    for el in root.iter():
                        if el.tag.rsplit("}", 1)[-1] == tag:
                            el.text = ""
                data = xml_bytes(root)
            elif name == "docProps/app.xml":
                root = etree_from_bytes(data)
                normalized = {
                    "Application": "Microsoft Office Word",
                    "Template": "Normal.dotm",
                    "TotalTime": "0",
                    "Company": "",
                    "Manager": "",
                }
                for tag in ("Pages", "Words", "Characters", "CharactersWithSpaces", "Lines", "Paragraphs"):
                    for el in root.iter():
                        if el.tag.rsplit("}", 1)[-1] == tag:
                            el.text = "0"
                for el in root.iter():
                    local = el.tag.rsplit("}", 1)[-1]
                    if local in normalized:
                        el.text = normalized[local]
                data = xml_bytes(root)
            elif name.startswith("word/") and name.endswith(".xml"):
                root = etree_from_bytes(data)
                changed = False
                for element in list(root.iter()):
                    namespace = element.tag[1:].split("}", 1)[0] if element.tag.startswith("{") else ""
                    if (
                        element.tag.rsplit("}", 1)[-1]
                        not in {"OLEObject", "attachedTemplate", "docVars"}
                        and namespace != "http://www.wps.cn/officeDocument/2013/wpsCustomData"
                    ):
                        continue
                    parent = element.getparent()
                    if parent is not None:
                        parent.remove(element)
                        changed = True
                if changed:
                    data = xml_bytes(root)
            elif name.endswith(".rels"):
                root = etree_from_bytes(data)
                for rel in list(root):
                    target = rel.get("Target", "")
                    if (
                        "embedding" in target
                        or target.startswith("../customXml")
                        or target.endswith("docProps/custom.xml")
                        or is_external_local_file_target(
                            target, rel.get("TargetMode", "")
                        )
                    ):
                        root.remove(rel)
                data = xml_bytes(root)
            elif name == "[Content_Types].xml":
                root = etree_from_bytes(data)
                for override in list(root):
                    part = override.get("PartName", "")
                    if (
                        "/word/embeddings/" in part
                        or "/customXml/" in part
                        or part == "/docProps/custom.xml"
                    ):
                        root.remove(override)
                data = xml_bytes(root)
            dst.writestr(info, data)
    for attempt in range(6):
        try:
            temp.replace(path)
            break
        except PermissionError:
            if attempt == 5:
                raise
            # Antivirus/indexing can briefly lock a newly written DOCX on
            # Windows. Keep replacement atomic and retry only this transient.
            time.sleep(0.05 * (attempt + 1))


def etree_from_bytes(data):
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    return etree.fromstring(data, parser=parser)


def xml_bytes(root):
    from lxml import etree
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def build(args):
    input_path = args.input.resolve()
    if not input_path.is_file() or input_path.suffix.lower() != ".json":
        raise FileNotFoundError(f"论文输入不存在或不是 .json: {input_path}")
    contest_config = load_contest_config(input_path.parent, args.contest_config)
    edition = contest_edition(contest_config)
    data = read_input(input_path)
    template = args.template.resolve()
    output = args.output.resolve()
    if not template.is_file() or template.suffix.lower() != ".docx":
        raise FileNotFoundError(f"Word 模板不存在或不是 .docx: {template}")
    if output.suffix.lower() != ".docx":
        raise ValueError(f"输出文件必须是 .docx: {output}")
    if output == template:
        raise ValueError("不得覆盖原始 Word 模板；请指定不同的 --output")
    manifest_out = args.manifest_out.resolve() if args.manifest_out else output.with_suffix(".chapters.json")
    if manifest_out.suffix.lower() != ".json":
        raise ValueError(f"--manifest-out 必须是 .json: {manifest_out}")
    protected = {input_path, template, output}
    if args.contest_config:
        protected.add(args.contest_config.resolve())
    if manifest_out in protected:
        raise ValueError("--manifest-out 不得覆盖输入、模板、配置或 DOCX 输出")
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document(str(template))
    section = clear_body_keep_final_section(document)
    configure_section(section)
    configure_header_footer(section)
    configure_styles(document)
    enable_update_fields(document)

    # Drafting DOCX front matter stays anonymous.  The formal 2026 upload is a
    # PDF and must receive the official identity cover in the final template.
    for text, size in [
        ("中国研究生创新实践系列大赛", 18),
        (f"“华为杯”第{edition}届中国研究生", 22),
        ("数学建模竞赛", 22),
    ]:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        r = p.add_run(text)
        set_run_font(r, "华文新魏", size, bold=True)

    title_p = document.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_p.paragraph_format.first_line_indent = Inches(0)
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(6)
    title_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    add_label_run(title_p, "题 目：")
    title_run = title_p.add_run(str(data["title"]).strip())
    set_run_font(title_run, "黑体", 16)

    abs_heading = document.add_paragraph()
    abs_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    abs_heading.paragraph_format.first_line_indent = Inches(0)
    abs_heading.paragraph_format.space_before = Pt(6)
    abs_heading.paragraph_format.space_after = Pt(6)
    abs_heading.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    add_label_run(abs_heading, "摘 要：")
    abstract_path = resolve_tex_path(input_path.parent.resolve(), str(data["abstract_tex_path"]), "abstract_tex_path")
    add_tex_content(document, read_tex_source(abstract_path, input_path.parent.resolve()), input_path.parent.resolve(), skip_first_section=False)

    keywords = document.add_paragraph()
    keywords.alignment = WD_ALIGN_PARAGRAPH.LEFT
    keywords.paragraph_format.first_line_indent = Inches(1 / 3)
    keywords.paragraph_format.space_before = Pt(6)
    keywords.paragraph_format.space_after = Pt(0)
    keywords.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    add_label_run(keywords, "关键词：")
    add_body_run(keywords, "  ".join(str(k).strip() for k in data["keywords"]))

    # The official format requires the next page after the complete abstract
    # to begin the body; do not insert a table of contents here.
    document.add_page_break()

    chapter_manifest = []
    problem_number = 0
    general_number = 0
    for chapter in data["chapters"]:
        title = str(chapter["title"]).strip()
        role = str(chapter.get("role", "problem"))
        level = int(chapter.get("level", 1))
        if level == 1:
            general_number += 1
            rendered, rendered_title = add_heading(document, title, level, general_number, role=role)
            if role == "problem":
                problem_number += 1
        else:
            rendered, rendered_title = add_heading(document, title, level)
        if role in {"references", "appendix"}:
            rendered.paragraph_format.page_break_before = True
        content_path = resolve_tex_path(input_path.parent.resolve(), str(chapter["tex_path"]), f"章节 {chapter['chapter_id']}")
        content = read_tex_source(content_path, input_path.parent.resolve())
        if not content.strip():
            raise ValueError(f"章节内容为空: {title}")
        add_tex_content(document, content, input_path.parent.resolve(), appendix=role == "appendix")
        chapter_manifest.append({
            "chapter_id": chapter["chapter_id"],
            "source_tex": str(chapter["tex_path"]),
            "title": title,
            "rendered_title": rendered_title,
            "role": role,
            "level": level,
            "aliases": chapter.get("aliases", []),
            "content_chars": len(content),
        })

    document.core_properties.title = str(data["title"]).strip()
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.subject = ""
    document.core_properties.keywords = ""
    document.save(str(output))
    scrub_package(output)
    manifest_out.parent.mkdir(parents=True, exist_ok=True)
    manifest_out.write_text(json.dumps({
        "title": data["title"],
        "keywords": data["keywords"],
        "chapters": chapter_manifest,
        "body_start_rule": "first page after the abstract/keywords page break",
        "body_end_rule": "page before references or appendix",
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "manifest": str(manifest_out)}, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--manifest-out", type=Path)
    parser.add_argument("--contest-config", type=Path, help="根目录比赛配置.json；省略时自动向上查找")
    args = parser.parse_args()
    build(args)


if __name__ == "__main__":
    main()
