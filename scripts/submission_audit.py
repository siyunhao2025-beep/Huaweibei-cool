#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Conservative pre-submission audit for a Huawei Cup paper.

The 2026 AI rules are conditional on how AI was used. This script never
treats an absent disclosure as evidence of a violation unless the user
explicitly declares relevant AI use with --ai-used.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path


IDENTITY_PAT = re.compile(
    r"(?:学校|学院|实验室|参赛队号|队员姓名|指导教师|学号|邮箱)\s*[:：]|"
    r"\\(?:schoolname|baominghao|member[abc]|makeidentitycover)\b|"
    r"\b(?:school\s*name|team\s*(?:number|id)|member\s*name|"
    r"advisor\s*name|student\s*(?:id|number)|e-?mail)\b|"
    r"C:\\Users\\",
    re.IGNORECASE,
)
AI_MARKER = re.compile(
    r"人工智能工具|人工智能辅助|AI\s*(?:工具|辅助|生成)|"
    r"本程序及代码是在人工智能工具辅助下完成",
    re.IGNORECASE,
)
COVER_MARKERS = ("学校", "参赛队号", "队员姓名")
COVER_FIELDS = ("schoolname", "baominghao", "membera", "memberb", "memberc")
COVER_PLACEHOLDER_RE = re.compile(
    r"REPLACE_WITH_ACTUAL(?:_[A-Z0-9_]+)?|\b(?:TODO|FIXME|TBD)\b|"
    r"待填|待补|待替换|请填写|请替换|占位",
    re.IGNORECASE,
)
SUBMISSION_SCHEMA = "submission-audit-v1"
SOURCE_RENDER_DPI = 72
SOURCE_INPUT_SUFFIXES = {
    ".tex", ".bib", ".bst", ".cls", ".sty", ".cfg", ".def",
    ".png", ".jpg", ".jpeg", ".pdf", ".eps", ".svg",
    ".csv", ".tsv", ".xlsx", ".xls", ".mat", ".npy", ".npz",
}
MANIFEST_REQUIRED_KEYS = {
    "title", "keywords", "abstract_tex_path", "appendix_pseudocode", "chapters",
}


def _resolved_tex_input(root: Path, raw: str) -> str:
    value = str(raw).strip().replace("\\", "/")
    if not value.lower().endswith(".tex"):
        value += ".tex"
    return os.path.normcase(str((root / value).resolve()))


def _source_tex_inputs(source: Path) -> list[str]:
    from audit_tex import INPUT_RE, remove_comments

    cleaned = remove_comments(source.read_text(encoding="utf-8-sig"))
    return [_resolved_tex_input(source.parent, raw) for raw in INPUT_RE.findall(cleaned)]


def _manifest_tex_inputs(path: Path, data: dict) -> list[str] | None:
    abstract_path = data.get("abstract_tex_path")
    chapters = data.get("chapters")
    if not isinstance(abstract_path, str) or not isinstance(chapters, list):
        return None
    raw_inputs = [abstract_path]
    for chapter in chapters:
        if not isinstance(chapter, dict) or not isinstance(chapter.get("tex_path"), str):
            return None
        raw_inputs.append(chapter["tex_path"])
    return [_resolved_tex_input(path.parent, raw) for raw in raw_inputs]


def resolve_source_manifest(source: Path, explicit: Path | None = None) -> dict:
    """Bind one validated paper manifest to the exact top-level TeX input chain."""
    source = source.resolve()
    same_dir = sorted(source.parent.glob("*.json"))
    conventional_names = {
        "manifest.json", "paper_manifest.json", "论文输入.json",
        f"{source.stem}.manifest.json",
    }
    parent_dir = [
        source.parent.parent / name
        for name in sorted(conventional_names)
        if (source.parent.parent / name).is_file()
    ]
    candidates = {item.resolve() for item in [*same_dir, *parent_dir] if item.is_file()}
    explicit_path = explicit.resolve() if explicit else None
    if explicit_path is not None:
        if not explicit_path.is_file():
            return {
                "status": "invalid", "path": explicit_path, "manifest": None,
                "error": "显式 manifest 不存在",
            }
        candidates = {explicit_path}

    try:
        source_inputs = _source_tex_inputs(source)
    except (OSError, UnicodeError) as exc:
        return {
            "status": "invalid", "path": explicit_path, "manifest": None,
            "error": f"TeX 主稿输入链无法读取: {exc}",
        }

    matching: list[tuple[Path, dict]] = []
    manifest_like: list[Path] = []
    invalid: list[tuple[Path, str]] = []
    for candidate in sorted(candidates):
        try:
            raw = json.loads(candidate.read_text(encoding="utf-8-sig"))
        except (OSError, UnicodeError, ValueError) as exc:
            if candidate == explicit_path:
                invalid.append((candidate, f"manifest JSON 无法解析: {exc}"))
            continue
        if not isinstance(raw, dict):
            if candidate == explicit_path:
                invalid.append((candidate, "manifest 顶层必须是 JSON 对象"))
            continue
        if candidate == explicit_path and not MANIFEST_REQUIRED_KEYS <= set(raw):
            missing = sorted(MANIFEST_REQUIRED_KEYS - set(raw))
            invalid.append((candidate, f"manifest 缺少字段: {', '.join(missing)}"))
            continue
        if not {"abstract_tex_path", "chapters"} <= set(raw):
            continue
        manifest_like.append(candidate)
        expected_inputs = _manifest_tex_inputs(candidate, raw)
        if expected_inputs != source_inputs:
            if candidate == explicit_path:
                invalid.append((candidate, "manifest 输入顺序与 TeX 主稿不一致"))
            continue
        try:
            from build_latex import load_manifest

            validated = load_manifest(candidate)
        except (OSError, UnicodeError, ValueError) as exc:
            invalid.append((candidate, str(exc)))
            continue
        matching.append((candidate, validated))

    if invalid:
        path, error = invalid[0]
        return {
            "status": "invalid", "path": path, "manifest": None,
            "source_inputs": source_inputs, "error": error,
        }
    if len(matching) > 1:
        return {
            "status": "ambiguous", "path": None, "manifest": None,
            "source_inputs": source_inputs,
            "candidates": [str(path) for path, _ in matching],
            "error": "多个 manifest 与同一 TeX 输入链匹配",
        }
    if matching:
        path, manifest = matching[0]
        if explicit_path is not None and path != explicit_path:
            return {
                "status": "invalid", "path": explicit_path, "manifest": None,
                "source_inputs": source_inputs,
                "error": "显式 manifest 未与 TeX 主稿输入链绑定",
            }
        return {
            "status": "bound", "path": path, "manifest": manifest,
            "source_inputs": source_inputs,
        }
    if explicit_path is not None:
        return {
            "status": "invalid", "path": explicit_path, "manifest": None,
            "source_inputs": source_inputs,
            "error": "显式 manifest 未与 TeX 主稿输入链绑定",
        }
    same_dir_manifests = [path for path in manifest_like if path.parent == source.parent]
    if same_dir_manifests:
        return {
            "status": "unbound", "path": same_dir_manifests[0], "manifest": None,
            "source_inputs": source_inputs,
            "error": "主稿同目录存在 manifest，但其输入链与当前 TeX 不一致",
        }
    return {
        "status": "not_found", "path": None, "manifest": None,
        "source_inputs": source_inputs,
    }


def canonical_sha256(value) -> str:
    data = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(data).hexdigest()


def file_sha256(path: Path | None) -> str | None:
    return hashlib.sha256(path.read_bytes()).hexdigest() if path and path.is_file() else None


def tree_sha256(path: Path | None) -> str | None:
    if path is None or not path.is_dir():
        return None
    files = sorted(item for item in path.rglob("*") if item.is_file())
    if not files:
        return None
    digest = hashlib.sha256()
    for item in files:
        digest.update(item.relative_to(path).as_posix().encode("utf-8"))
        digest.update(b"\0")
        digest.update(item.read_bytes())
        digest.update(b"\0")
    return digest.hexdigest()


def source_tree_sha256(source: Path | None) -> str | None:
    """Hash manuscript inputs while excluding generated reports/build products."""
    if source is None or not source.is_file():
        return None
    source = source.resolve()
    root = source.parent
    files = []
    for item in root.rglob("*"):
        if not item.is_file() or item.suffix.lower() not in SOURCE_INPUT_SUFFIXES:
            continue
        if item.resolve() == source.with_suffix(".pdf").resolve():
            continue
        files.append(item)
    if source not in files:
        files.append(source)
    manifest_binding = resolve_source_manifest(source)
    bound_manifest = manifest_binding.get("path") if manifest_binding.get("status") == "bound" else None
    if isinstance(bound_manifest, Path) and bound_manifest not in files:
        files.append(bound_manifest)
    digest = hashlib.sha256()
    for item in sorted(set(files)):
        digest.update(Path(os.path.relpath(item, root)).as_posix().encode("utf-8"))
        digest.update(b"\0")
        digest.update(item.read_bytes())
        digest.update(b"\0")
    return digest.hexdigest()


def relative_path(path: Path | None, base: Path) -> str | None:
    if path is None:
        return None
    try:
        return Path(os.path.relpath(path, base)).as_posix()
    except ValueError:
        return None


def build_submission_provenance(args, report_dir: Path, audit_payload: dict) -> dict:
    paper = Path(args.paper).resolve()
    source = Path(args.source).resolve() if args.source else None
    attachments = Path(args.attachments).resolve() if args.attachments else None
    ai_file = Path(args.ai_file).resolve() if args.ai_file else None
    contract = audit_payload.get("manifest_pdf_contract") or {}
    binding = contract.get("binding") or {}
    manifest = None
    if source is not None and binding.get("status") == "bound" and binding.get("manifest_path"):
        candidate = (source.parent / binding["manifest_path"]).resolve()
        if candidate.is_file():
            manifest = candidate
    audit_view = {
        key: audit_payload[key]
        for key in ("passed", "pages", "ai_used", "cover_policy", "checks", "warnings")
    }
    paper_fingerprints = pdf_fingerprints(paper) if paper.suffix.lower() == ".pdf" else {}
    payload = {
        "schema": SUBMISSION_SCHEMA,
        "generator_sha256": file_sha256(Path(__file__).resolve()),
        "paper_path": relative_path(paper, report_dir),
        "paper_sha256": file_sha256(paper),
        "paper_page_count": paper_fingerprints.get("page_count"),
        "paper_page_sizes": paper_fingerprints.get("page_sizes"),
        "paper_text_sha256": paper_fingerprints.get("text_sha256"),
        "paper_render_sha256": paper_fingerprints.get("render_sha256"),
        "source_path": relative_path(source, report_dir),
        "source_tree_sha256": source_tree_sha256(source),
        "manifest_path": relative_path(manifest, report_dir),
        "manifest_sha256": file_sha256(manifest),
        "attachments_path": relative_path(attachments, report_dir),
        "attachments_sha256": tree_sha256(attachments),
        "ai_file_path": relative_path(ai_file, report_dir),
        "ai_file_sha256": file_sha256(ai_file),
        "parameters": {
            "max_total_pages": args.max_total_pages,
            "ai_used": args.ai_used,
            "cover_policy": args.cover_policy,
        },
        "audit_sha256": canonical_sha256(audit_view),
    }
    return {**payload, "fingerprint": canonical_sha256(payload)}


def pdf_text_sha256(path: Path) -> str | None:
    """Stable page-wise text fingerprint for a PDF, independent of metadata."""
    pages = extract_pdf_pages(path)
    if not pages:
        return None
    normalized = [
        re.sub(r"\s+", "", unicodedata.normalize("NFKC", page))
        for page in pages
    ]
    return canonical_sha256(normalized)


def pdf_fingerprints(path: Path) -> dict:
    """Fingerprint page count, boxes, text layer and fixed-DPI rendered pixels."""
    try:
        pymupdf = load_pymupdf()
        digest = hashlib.sha256()
        with pymupdf.open(path) as document:
            sizes = []
            for page in document:
                sizes.append([round(float(page.rect.width), 3), round(float(page.rect.height), 3)])
                pixmap = page.get_pixmap(
                    dpi=SOURCE_RENDER_DPI,
                    colorspace=pymupdf.csGRAY,
                    alpha=False,
                )
                digest.update(f"{pixmap.width}x{pixmap.height}\0".encode("ascii"))
                digest.update(pixmap.samples)
            page_count = document.page_count
    except Exception:
        return {}
    return {
        "page_count": page_count,
        "page_sizes": sizes,
        "text_sha256": pdf_text_sha256(path),
        "render_sha256": digest.hexdigest(),
    }


def compile_tex_fingerprints(source: Path) -> tuple[dict | None, str]:
    """Compile a TeX entry in an isolated output directory and fingerprint it."""
    latexmk = shutil.which("latexmk")
    xelatex = shutil.which("xelatex")
    if latexmk is None and xelatex is None:
        return None, "未找到 latexmk/xelatex，无法证明最终 PDF 与 TeX 源同源"
    try:
        with tempfile.TemporaryDirectory(prefix="gmcm-source-build-") as output_dir:
            run = None
            if latexmk is not None:
                run = subprocess.run(
                    [
                        latexmk,
                        "-xelatex",
                        "-interaction=nonstopmode",
                        "-halt-on-error",
                        "-file-line-error",
                        f"-outdir={output_dir}",
                        source.name,
                    ],
                    cwd=source.parent,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=120,
                )
            if (run is None or run.returncode != 0) and xelatex is not None:
                command = [
                    xelatex,
                    "-interaction=nonstopmode",
                    "-halt-on-error",
                    "-file-line-error",
                    f"-output-directory={output_dir}",
                    source.name,
                ]
                for _ in range(2):
                    run = subprocess.run(
                        command,
                        cwd=source.parent,
                        capture_output=True,
                        text=True,
                        encoding="utf-8",
                        errors="replace",
                        timeout=120,
                    )
                    if run.returncode != 0:
                        break
            compiled = Path(output_dir) / f"{source.stem}.pdf"
            if run is None or run.returncode != 0 or not compiled.is_file():
                tail = (run.stdout + "\n" + run.stderr).strip().splitlines()[-3:]
                return None, "TeX 现场编译失败：" + " | ".join(tail)
            fingerprints = pdf_fingerprints(compiled)
            if not fingerprints or fingerprints.get("text_sha256") is None:
                return None, "TeX 现场编译产物无可读内容"
            return fingerprints, "TeX 现场编译成功"
    except (OSError, subprocess.TimeoutExpired) as exc:
        return None, f"TeX 现场编译失败：{exc}"


def load_pymupdf():
    try:
        import pymupdf  # type: ignore
    except ImportError:
        import fitz as pymupdf  # type: ignore
    return pymupdf


def count_pdf_pages(pdf: Path) -> int | None:
    try:
        import pypdf  # type: ignore
        with pdf.open("rb") as handle:
            return len(pypdf.PdfReader(handle).pages)
    except Exception:
        pass
    try:
        pymupdf = load_pymupdf()
        with pymupdf.open(pdf) as document:
            return document.page_count
    except Exception:
        return None


def strip_tex_comments(text: str) -> str:
    cleaned = []
    for line in text.splitlines():
        cut = len(line)
        for index, char in enumerate(line):
            if char != "%":
                continue
            backslashes = 0
            cursor = index - 1
            while cursor >= 0 and line[cursor] == "\\":
                backslashes += 1
                cursor -= 1
            if backslashes % 2 == 0:
                cut = index
                break
        cleaned.append(line[:cut])
    return "\n".join(cleaned)


def extract_pdf_pages(path: Path) -> list[str]:
    try:
        pymupdf = load_pymupdf()
        with pymupdf.open(path) as document:
            return [page.get_text() for page in document]
    except Exception:
        return []


def strip_tex_cover(text: str) -> str:
    """Remove only the officially permitted page-0 identity commands."""
    clean = strip_tex_comments(text)
    for command in COVER_FIELDS:
        clean = re.sub(rf"\\{command}\s*\{{[^{{}}]*\}}", "", clean, flags=re.S)
    return re.sub(r"\\makeidentitycover\b", "", clean)


def tex_cover_present(text: str) -> bool:
    clean = strip_tex_comments(text)
    return bool(re.search(r"\\makeidentitycover\b", clean)) and all(
        re.search(rf"\\{command}\s*\{{", clean) for command in COVER_FIELDS
    )


def tex_value_has_visible_content(value: str) -> bool:
    plain = value
    previous = None
    while previous != plain:
        previous = plain
        plain = re.sub(
            r"\\(?:phantom|hphantom|vphantom|smash|rlap|llap|clap|hbox|mbox|makebox)"
            r"(?:\s*\[[^]]*\])?\s*\{[^{}]*\}",
            "",
            plain,
            flags=re.IGNORECASE,
        )
        plain = re.sub(
            r"\\(?:hspace|vspace)\*?\s*\{[^{}]*\}|"
            r"\\kern\s*[-+]?\d*\.?\d+(?:pt|em|ex|mu|mm|cm|in|pc|bp|dd|cc|sp)\b",
            "",
            plain,
            flags=re.IGNORECASE,
        )
    plain = re.sub(
        r"\\(?:quad|qquad|enspace|thinspace|medspace|thickspace|"
        r"negthinspace|negmedspace|negthickspace|strut|null|hfill|vfill|relax)\b|\\[,;!:]",
        "",
        plain,
        flags=re.IGNORECASE,
    )
    plain = re.sub(r"[\s${}~]+", "", plain)
    return bool(plain)


def cover_value_is_valid(value: str) -> bool:
    """Reject blank, sentinel and line-like values in an identity field."""
    compact = re.sub(r"\s+", "", unicodedata.normalize("NFKC", value or ""))
    if not compact or COVER_PLACEHOLDER_RE.search(compact):
        return False
    if re.fullmatch(r"[_\-\u2010-\u2015─－＿.]+", compact):
        return False
    if re.fullmatch(r"x{1,12}", compact, re.IGNORECASE):
        return False
    return True


def tex_cover_fields_complete(text: str) -> bool:
    """Require all five official TeX cover values to be visible and non-placeholder."""
    clean = strip_tex_comments(text)
    for command in COVER_FIELDS:
        match = re.search(rf"\\{command}\s*\{{([^{{}}]*)\}}", clean, re.S)
        if not match or not tex_value_has_visible_content(match.group(1)):
            return False
        if not cover_value_is_valid(match.group(1)):
            return False
    return True


def pdf_cover_fields_complete(path: Path) -> bool:
    """Check value zones on the fixed official page-0 layout.

    The logo/content identity of the official cover remains a manual check. This
    positional test only prevents an untouched blank cover from passing formal
    submission audit.
    """
    try:
        pymupdf = load_pymupdf()
        with pymupdf.open(path) as document:
            if document.page_count < 1:
                return False
            words = document[0].get_text("words")
    except Exception:
        return False

    value_zones = (
        (160.0, 345.0, 390.0),
        (170.0, 390.0, 430.0),
        (215.0, 430.0, 466.0),
        (215.0, 466.0, 505.0),
        (215.0, 505.0, 545.0),
    )
    for min_x, min_y, max_y in value_zones:
        values = []
        for x0, y0, _x1, y1, value, *_rest in words:
            text = str(value).strip()
            center_y = (float(y0) + float(y1)) / 2
            if float(x0) >= min_x and min_y <= center_y <= max_y and text:
                values.append(text)
        if not cover_value_is_valid(" ".join(values)):
            return False
    return True


def extract_text(path: Path) -> str:
    """Extract enough text for conservative content checks."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            pymupdf = load_pymupdf()
            with pymupdf.open(path) as document:
                return "\n".join(page.get_text() for page in document)
        except Exception:
            return ""
    if suffix in {".tex", ".txt", ".md"}:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except (OSError, UnicodeError):
            return ""
        return strip_tex_comments(text) if suffix == ".tex" else text
    if suffix == ".docx":
        try:
            from docx import Document  # type: ignore
            document = Document(str(path))
            blocks = [paragraph.text for paragraph in document.paragraphs]
            blocks.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            for section in document.sections:
                blocks.extend(paragraph.text for paragraph in section.header.paragraphs)
                blocks.extend(paragraph.text for paragraph in section.footer.paragraphs)
            return "\n".join(blocks)
        except Exception:
            return ""
    return ""


def audit_manifest_bound_pdf(
        paper: Path, source: Path, explicit_manifest: Path | None = None,
        *, audit_front_matter: bool = True) -> dict:
    """Run rendered front/terminal page gates against one source-bound manifest."""
    resolution = resolve_source_manifest(source, explicit_manifest)
    manifest_path = resolution.get("path")
    binding = {
        "status": resolution["status"],
        "manifest_path": (
            relative_path(manifest_path, source.resolve().parent)
            if isinstance(manifest_path, Path) else None
        ),
        "manifest_sha256": (
            file_sha256(manifest_path) if isinstance(manifest_path, Path) else None
        ),
        "source_input_count": len(resolution.get("source_inputs", [])),
        "error": resolution.get("error"),
    }
    if resolution["status"] == "not_found":
        return {
            "applicable": False,
            "ok": True,
            "binding": binding,
            "abstract_front_matter": None,
            "terminal_chapters": None,
        }
    if resolution["status"] != "bound":
        return {
            "applicable": True,
            "ok": False,
            "binding": binding,
            "abstract_front_matter": None,
            "terminal_chapters": None,
        }

    try:
        from audit_paper import (
            abstract_front_matter_audit,
            extract_page_records,
            terminal_chapter_page_audit,
        )

        records = extract_page_records(paper)
        manifest = resolution["manifest"]
        front_matter = (
            abstract_front_matter_audit(records, manifest)
            if audit_front_matter else None
        )
        terminal_pages = terminal_chapter_page_audit(
            records, manifest, manifest_supplied=True
        )
    except Exception as exc:
        binding["error"] = f"成品 PDF 章节审计无法执行: {exc}"
        return {
            "applicable": True,
            "ok": False,
            "binding": binding,
            "abstract_front_matter": None,
            "terminal_chapters": None,
        }

    front_ok = (
        True if front_matter is None else
        bool(front_matter["details"].get("verifiable")) and not front_matter["failures"]
    )
    terminal_ok = (
        bool(terminal_pages["details"].get("verifiable"))
        and not terminal_pages["failures"]
    )
    return {
        "applicable": True,
        "ok": front_ok and terminal_ok,
        "binding": binding,
        "abstract_front_matter": front_matter,
        "terminal_chapters": terminal_pages,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="华为杯提交前审计")
    parser.add_argument("--paper", required=True, help="论文 PDF / tex / docx 路径")
    parser.add_argument(
        "--source",
        help="最终 PDF 对应的 TeX 主稿；P6 会现场编译并校验同源",
    )
    parser.add_argument(
        "--manifest",
        help="可选的原始论文 manifest；必须与 TeX 输入链一致，路径与哈希由 P6 绑定",
    )
    parser.add_argument(
        "--max-total-pages", "--body-gate", dest="max_total_pages", type=int,
        help="总页数上限（仅在当届官方明确时使用；--body-gate 为兼容别名）",
    )
    parser.add_argument(
        "--ai-used", choices=("unknown", "none", "writing", "analysis", "programming", "mixed"),
        default="unknown", help="实际 AI 使用方式；默认 unknown 不将未披露误判为违规",
    )
    parser.add_argument("--ai-file", help="AI 使用/标注说明文件（声明使用 AI 时的审计证据）")
    parser.add_argument("--attachments", help="提交附件目录")
    parser.add_argument(
        "--cover-policy", choices=("required", "forbidden"), default="required",
        help="2026 正式提交默认 required；纯匿名内部审阅稿使用 forbidden",
    )
    parser.add_argument("--report", help="写入绑定当前输入的结构化 JSON 审计报告")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)

    paper = Path(args.paper)
    if not paper.is_file():
        print(f"[FAIL] 论文不存在: {paper}", file=sys.stderr)
        return 1

    checks: list[dict] = []
    warnings: list[str] = []

    def check(ok: bool, desc: str, **details) -> None:
        checks.append({"ok": bool(ok), "desc": desc, **details})

    pages = None
    if paper.suffix.lower() == ".pdf":
        pages = count_pdf_pages(paper)
        check(pages is not None, "PDF 页数可读取")
        if pages is not None:
            check(True, f"PDF 共 {pages} 页")
            if args.max_total_pages is not None:
                check(
                    pages <= args.max_total_pages,
                    f"总页数 {pages} ≤ 官方上限 {args.max_total_pages}",
                )
    else:
        warnings.append("非 PDF 文件未统计页数；请以最终 PDF 人工核对摘要页、页码和版式。")

    suffix = paper.suffix.lower()
    supported_suffix = suffix in {".pdf", ".tex", ".txt", ".md", ".docx"}
    check(supported_suffix, f"文件类型受支持: {suffix or '(无扩展名)'}")
    text = extract_text(paper) if supported_suffix else ""
    check(bool(text.strip()), "正文文本可读取，匿名与披露检查可执行")
    manifest_pdf_contract = None
    if args.source:
        source = Path(args.source)
        source_ok = source.is_file() and source.suffix.lower() == ".tex"
        check(source_ok, "已绑定可读的 TeX 主稿")
        if source_ok and suffix == ".pdf":
            compiled_signature, compile_detail = compile_tex_fingerprints(source.resolve())
            final_signature = pdf_fingerprints(paper.resolve())
            check(
                compiled_signature is not None
                and final_signature is not None
                and compiled_signature == final_signature,
                "最终 PDF 与当前 TeX 主稿现场编译结果同源"
                "（页数、页面尺寸、文本与固定 DPI 像素一致）"
                + (f"（{compile_detail}）" if compile_detail else ""),
            )
            manifest_pdf_contract = audit_manifest_bound_pdf(
                paper.resolve(),
                source.resolve(),
                Path(args.manifest) if args.manifest else None,
                audit_front_matter=args.cover_policy == "required",
            )
            binding = manifest_pdf_contract["binding"]
            if not manifest_pdf_contract["applicable"]:
                manifest_required = bool(args.report) or args.cover_policy == "required"
                check(
                    not manifest_required,
                    (
                        "正式提交审计必须绑定与 TeX 输入链匹配的原始 manifest"
                        if manifest_required else
                        "非正式独立调用未发现原始 manifest；成品摘要与终端章节分页门禁不适用"
                    ),
                    code="paper_manifest_binding",
                    evidence=binding,
                )
            else:
                bound = binding["status"] == "bound"
                check(
                    bound,
                    "原始 manifest 已唯一绑定 TeX 输入链，且路径与 SHA-256 已纳入提交审计",
                    code="paper_manifest_binding",
                    evidence=binding,
                )
                if bound:
                    front_matter = manifest_pdf_contract["abstract_front_matter"]
                    if front_matter is not None:
                        front_ok = (
                            bool(front_matter["details"].get("verifiable"))
                            and not front_matter["failures"]
                        )
                        check(
                            front_ok,
                            "正式成品 PDF 的题目、摘要、关键词同处物理第 2 页，正文从第 3 页开始，字号与题目行数合规",
                            code="abstract_front_matter",
                            evidence={"binding": binding, "result": front_matter},
                        )
                    terminal = manifest_pdf_contract["terminal_chapters"]
                    terminal_ok = (
                        bool(terminal["details"].get("verifiable"))
                        and not terminal["failures"]
                    )
                    check(
                        terminal_ok,
                        "manifest 恰声明一个参考文献章节；已声明附录，且参考文献与各附录均从物理页首开始",
                        code="terminal_chapter_pages",
                        evidence={"binding": binding, "result": terminal},
                    )
        else:
            check(False, "同源校验要求最终稿为 PDF，源稿为 TeX")
    elif args.manifest:
        check(False, "--manifest 必须与 --source TeX 主稿同时使用")
    cover_present = False
    cover_fields_complete: bool | None = None
    anonymous_text = text
    can_separate_cover = False
    if suffix == ".pdf":
        page_texts = extract_pdf_pages(paper)
        if page_texts:
            cover_compact = re.sub(r"\s+", "", page_texts[0])
            cover_present = all(marker in cover_compact for marker in COVER_MARKERS)
            anonymous_text = "\n".join(page_texts[1:])
            can_separate_cover = len(page_texts) >= 2
            cover_fields_complete = pdf_cover_fields_complete(paper)
    elif suffix == ".tex":
        cover_present = tex_cover_present(text)
        cover_fields_complete = tex_cover_fields_complete(text)
        anonymous_text = strip_tex_cover(text)
        can_separate_cover = True
    elif suffix == ".docx":
        compact = re.sub(r"\s+", "", text)
        cover_present = all(marker in compact for marker in COVER_MARKERS)
        warnings.append("DOCX 无法可靠按页分离封皮；封皮后的匿名性须以最终 PDF 再审计。")

    if args.cover_policy == "required":
        check(cover_present, "2026 正式提交含第 0 页官方参赛信息封皮")
        if cover_fields_complete is not None:
            check(
                cover_fields_complete,
                "第 0 页学校、参赛队号和三名队员姓名均已填写且不是占位文本",
            )
        else:
            warnings.append("当前格式无法可靠核验封皮字段是否填写；请以最终 PDF 再审计。")
        warnings.append("人工核对：封皮四个官方 logo 未删除、替换或变形。")
        if can_separate_cover and anonymous_text:
            hits = sorted(set(IDENTITY_PAT.findall(anonymous_text)))
            check(not hits, f"封皮之后无身份标志（检出 {len(hits)} 处疑似项：{hits[:10]}）")
        elif not can_separate_cover:
            warnings.append("未能可靠分离封皮与正文，封皮后的身份信息需人工复核。")
    else:
        check(not cover_present, "内部纯匿名审阅稿不含参赛信息封皮")
        if text:
            hits = sorted(set(IDENTITY_PAT.findall(text)))
            check(not hits, f"纯匿名稿无身份标志（检出 {len(hits)} 处疑似项：{hits[:10]}）")
        else:
            warnings.append("未能抽取正文文本，身份信息需人工复核。")

    # Public sources and borrowed programs need formal references. A template
    # need not manufacture a citation, so this is a review reminder, not a
    # false automated failure.
    check(True, "引用格式：如使用公开资料或程序，正文应按 [n] 引用并列完整参考文献。")

    disclosure_exists = bool(
        args.ai_file
        and Path(args.ai_file).is_file()
        and Path(args.ai_file).stat().st_size > 0
    )
    marker_exists = bool(text and AI_MARKER.search(text))
    if args.ai_used == "none":
        check(
            not disclosure_exists and not marker_exists,
            "AI 使用声明为 none，且未发现相互矛盾的 AI 说明文件或论文内标注。",
        )
    elif args.ai_used == "unknown":
        warnings.append("尚未声明是否使用 AI；提交前须人工确认 --ai-used=none 或相应使用方式。")
    else:
        check(
            disclosure_exists or marker_exists,
            "已声明使用 AI，且存在 AI 使用/标注证据（说明文件或论文内标注）。",
        )
        if args.ai_used in {"analysis", "mixed"}:
            warnings.append("人工核对：AI 辅助数据分析的结果前后须标明工具、版本、开发者和发布日期。")
        if args.ai_used in {"programming", "mixed"}:
            warnings.append("人工核对：AI 辅助代码开头须有规定的工具信息注释。")
        if args.ai_used in {"writing", "mixed"}:
            warnings.append("人工核对：最终文字须经队伍理解并用自己的语言表述。")

    if args.ai_file:
        check(disclosure_exists, f"AI 说明文件存在: {Path(args.ai_file).name}")

    if args.attachments:
        attachment_dir = Path(args.attachments)
        files = [item for item in attachment_dir.rglob("*") if item.is_file()] if attachment_dir.is_dir() else []
        check(attachment_dir.is_dir() and bool(files), f"附件目录含 {len(files)} 个文件")

    passed = all(item["ok"] for item in checks)
    payload = {
        "paper": str(paper),
        "passed": passed,
        "pages": pages,
        "ai_used": args.ai_used,
        "cover_policy": args.cover_policy,
        "checks": checks,
        "warnings": warnings,
        "manifest_pdf_contract": manifest_pdf_contract,
    }
    if args.report:
        report_path = Path(args.report).resolve()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        provenance = build_submission_provenance(args, report_path.parent, payload)
        payload = {
            **payload,
            "schema": SUBMISSION_SCHEMA,
            "paper": provenance["paper_path"],
            "provenance": provenance,
        }
        report_path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    if args.json:
        print(json.dumps(payload, ensure_ascii=False, indent=2))
    else:
        print(f"=== 提交前审计：{paper.name} ===")
        for item in checks:
            print(f"  [{'x' if item['ok'] else ' '}] {item['desc']}")
        for warning in warnings:
            print(f"  [!] {warning}")
        print(f"\n结果: {'PASS' if passed else 'FAIL'}")
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
